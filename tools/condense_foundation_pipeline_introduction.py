from pathlib import Path
import shutil

from docx import Document


MANUSCRIPT = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/manuscript/"
    "manuscript_JTM_multifoundation_atlas.docx"
)
BACKUP = MANUSCRIPT.with_name(
    "manuscript_JTM_multifoundation_atlas.before_foundation_intro_condense.docx"
)
OLD_START = "We evaluated three released embedding pipelines"
NEW_TEXT = (
    "TITAN is a multimodal whole-slide model trained with visual self-supervision and "
    "vision-language alignment; its published Mass-340K pretraining corpus excluded TCGA, although "
    "TCGA was used for downstream evaluation [17]. Giga-SSL is a gigapixel self-supervised whole-slide "
    "representation whose official repository distributes 512-dimensional TCGA embeddings [18,19]; "
    "TCGA was used during Giga-SSL development, and direct overlap between its representation-learning "
    "images and the evaluated slides cannot be excluded. Prov-GigaPath combines a tile encoder with a "
    "long-context slide encoder trained on Providence health-system pathology data; we used the final "
    "768-dimensional slide layer from the public TCGA embedding dataset [20,21]. No weights were "
    "fine-tuned and no downstream molecular labels entered representation learning in this study. "
    "Giga-SSL is therefore label-held-out in the downstream cross-validation, but not necessarily "
    "image-unseen at the representation-learning stage."
)


def replace_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def main():
    shutil.copy2(MANUSCRIPT, BACKUP)
    document = Document(MANUSCRIPT)
    matches = [p for p in document.paragraphs if p.text.startswith(OLD_START)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one pipeline paragraph, found {len(matches)}")
    replace_paragraph_text(matches[0], NEW_TEXT)
    document.save(MANUSCRIPT)
    print(f"Updated: {MANUSCRIPT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
