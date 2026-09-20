#!/usr/bin/env python3
"""Remove the patient-overlap Venn from the current manuscript and renumber figures."""

from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
BACKUP = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.before_remove_figure1.docx"


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def paragraph_starting(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"Paragraph not found: {prefix}")


copy2(DOCX, BACKUP)
document = Document(DOCX)

cohort = paragraph_starting(document, "The released artifacts comprised")
cohort.text = cohort.text.replace(
    "; Figure 1 shows the pairwise and dataset-specific counts.", "."
)

caption = paragraph_starting(document, "Figure 1. Patient overlap among")
caption_index = next(
    index for index, paragraph in enumerate(document.paragraphs)
    if paragraph._p is caption._p
)
if caption_index == 0 or not document.paragraphs[caption_index - 1]._p.xpath(".//w:drawing"):
    raise RuntimeError("The expected Figure 1 drawing was not found before its caption")
drawing = document.paragraphs[caption_index - 1]
remove_paragraph(drawing)
remove_paragraph(caption)

figure_2 = paragraph_starting(
    document, "Figure 2. Provenance-stratified effect-threshold crossing rates."
)
figure_2.text = figure_2.text.replace("Figure 2.", "Figure 1.", 1)

figure_3 = paragraph_starting(document, "Figure 3. Partition sensitivity under")
figure_3.text = figure_3.text.replace("Figure 3.", "Figure 2.", 1)
figure_3.text = figure_3.text.replace("match Figure 2.", "match Figure 1.")

document.save(DOCX)
print(DOCX)
