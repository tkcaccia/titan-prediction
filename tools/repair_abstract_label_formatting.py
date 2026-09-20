#!/usr/bin/env python3
"""Restore bold-only abstract labels after a paragraph-level text edit."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
LABEL = "Methods."


document = Document(DOCX)
matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(LABEL)]
if len(matches) != 1:
    raise RuntimeError(f"Expected one abstract Methods paragraph, found {len(matches)}")

paragraph = matches[0]
body = paragraph.text[len(LABEL) :].lstrip()
for run in paragraph.runs:
    run.text = ""

label_run = paragraph.runs[0]
label_run.text = LABEL
label_run.bold = True
body_run = paragraph.add_run(" " + body)
body_run.bold = False

document.save(DOCX)
print(f"Repaired abstract label formatting in {DOCX}")
