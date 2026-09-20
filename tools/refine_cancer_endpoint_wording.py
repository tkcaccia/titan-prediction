#!/usr/bin/env python3
"""Refine the model-per-representation wording in the current manuscript."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
REPLACEMENTS = {
    "The pair was fitted separately to each available representation": (
        "A separate model was then fitted for each available representation"
    ),
    "It defines one cancer-specific prediction task; each representation is fitted "
    "separately, and no model pools cancer types.": (
        "It defines one cancer-specific task; a separate model is fitted for each "
        "representation, and no model pools cancer types."
    ),
}


document = Document(DOCX)
changed = 0

for paragraph in document.paragraphs:
    revised = paragraph.text
    for old, new in REPLACEMENTS.items():
        revised = revised.replace(old, new)
    if revised != paragraph.text:
        paragraph.runs[0].text = revised
        for run in paragraph.runs[1:]:
            run.text = ""
        changed += 1

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                revised = paragraph.text
                for old, new in REPLACEMENTS.items():
                    revised = revised.replace(old, new)
                if revised != paragraph.text:
                    paragraph.runs[0].text = revised
                    for run in paragraph.runs[1:]:
                        run.text = ""
                    changed += 1

if changed != 2:
    raise RuntimeError(f"Expected two wording changes, found {changed}")

document.save(DOCX)
print(f"Refined {changed} cancer-endpoint explanations in {DOCX}")
