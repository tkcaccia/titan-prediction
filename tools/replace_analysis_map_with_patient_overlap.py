#!/usr/bin/env python3

"""Replace the main analysis-map table and generic workflow with a patient Venn."""

from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path("/Users/stefano/Documents/Titan/titan-prediction")
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
BACKUP = ROOT / "manuscript" / (
    "manuscript_JTM_multifoundation_atlas.before_patient_overlap_venn.docx"
)
FIGURE = ROOT / "results" / "figures" / "Figure1_patient_overlap_venn.png"


def paragraph_starting(document, prefix):
    return next(p for p in document.paragraphs if p.text.startswith(prefix))


def replace_text(paragraph, text):
    for child in list(paragraph._p):
        if not child.tag.endswith("}pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


def remove_paragraph(paragraph):
    paragraph._p.getparent().remove(paragraph._p)


if not DOCX.exists():
    raise FileNotFoundError(DOCX)
if not FIGURE.exists():
    raise FileNotFoundError(FIGURE)
if not BACKUP.exists():
    shutil.copy2(DOCX, BACKUP)

doc = Document(DOCX)

# Remove the former analysis-map table and its caption.
analysis_caption = paragraph_starting(doc, "Table 1. Analysis map.")
remove_paragraph(analysis_caption)
analysis_table = doc.tables[0]
analysis_table._element.getparent().remove(analysis_table._element)

analysis_heading = paragraph_starting(doc, "Analysis map and evidence hierarchy")
replace_text(analysis_heading, "Cohort overlap and evidence hierarchy")
hierarchy = doc.add_paragraph(
    "The primary analysis is the matched three-representation benchmark. "
    "The supporting TITAN screen adds permutation/FDR qualification for one "
    "representation, while PathoFMPred records analysis and model metadata "
    "without creating new performance evidence. Results from these layers "
    "are not interchangeable.",
    style="Normal",
)
analysis_heading._p.addnext(hierarchy._p)

# Expand the cohort description with mutually exclusive overlap counts.
cohort = paragraph_starting(doc, "The released artifacts comprised 11,449 TITAN slides")
replace_text(
    cohort,
    "The released artifacts comprised 11,449 TITAN slides from 9,404 patients, "
    "11,427 Giga-SSL slides from 9,378 patients and 10,328 Prov-GigaPath slides "
    "from 8,393 patients. Their union contained 9,471 unique patients. Of these, "
    "8,241 (87.0%) occurred in all three datasets and formed the matched benchmark. "
    "Pairwise-only overlaps were 1,070 for TITAN and Giga-SSL, 92 for TITAN and "
    "Prov-GigaPath, and 60 for Giga-SSL and Prov-GigaPath; only 1, 7 and 0 patients "
    "were unique to the respective datasets.",
)
pooling = doc.add_paragraph(
    "Within each representation, eligible primary-tumour diagnostic slides were "
    "pooled by arithmetic mean before outcomes were joined; the patient, rather "
    "than the slide, was the unit of analysis and cross-validation. Patients with "
    "multiple slides remained one observation, and all slides from one patient "
    "stayed in the same fold.",
    style="Normal",
)
cohort._p.addnext(pooling._p)

# Replace the generic patient-workflow image with the patient-overlap Venn and
# place it immediately after the cohort and pooling description.
old_caption = paragraph_starting(doc, "Figure 1. Patient-first design.")
image_paragraph = old_caption._p.getprevious()
for child in list(image_paragraph):
    if not child.tag.endswith("}pPr"):
        image_paragraph.remove(child)
image_paragraph_wrapper = next(
    p for p in doc.paragraphs if p._p is image_paragraph
)
image_paragraph_wrapper.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = image_paragraph_wrapper.add_run()
shape = run.add_picture(str(FIGURE), width=Inches(6.35))
shape._inline.docPr.set(
    "descr",
    "Non-area-proportional Venn diagram of patient overlap among TITAN, "
    "Giga-SSL and Prov-GigaPath TCGA embedding datasets.",
)
shape._inline.docPr.set("title", "Patient overlap across embedding datasets")
replace_text(
    old_caption,
    "Figure 1. Patient overlap among the released TITAN, Giga-SSL and "
    "Prov-GigaPath TCGA embedding datasets. Values are mutually exclusive "
    "region counts. The 8,241-patient three-way intersection was used for the "
    "matched benchmark. Circle areas are schematic and are not proportional "
    "to cohort size.",
)
old_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

pooling._p.addnext(image_paragraph_wrapper._p)
image_paragraph_wrapper._p.addnext(old_caption._p)

# The provenance table is now the first numbered main-text table.
provenance_caption = paragraph_starting(
    doc, "Table 2. Provenance-stratified effect-threshold crossings."
)
replace_text(
    provenance_caption,
    provenance_caption.text.replace("Table 2.", "Table 1.", 1),
)

doc.save(DOCX)
print(DOCX)
