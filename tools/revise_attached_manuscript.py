from copy import deepcopy
from pathlib import Path
import re
import shutil
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path('/private/tmp/jtm_revision_20260916/accepted.docx')
OUTPUT = Path('/Users/stefano/Documents/Titan/titan-prediction/manuscript/manuscript_JTM_multifoundation_atlas_revised.docx')
FIGURE2 = Path('/Users/stefano/Documents/Titan/titan-prediction/figures/Figure2_biological_predictability_map.png')


def replace_paragraph(paragraph, text, bold_prefix=None):
    paragraph.clear()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def insert_after(paragraph, text):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph._p.pPr is not None:
        new_para._p.insert(0, deepcopy(paragraph._p.pPr))
    new_para.add_run(text)
    return new_para


def shift_citations(text, increment=2):
    def repl(match):
        inner = match.group(1)
        if not re.fullmatch(r'[0-9,;\-\s]+', inner):
            return match.group(0)
        def shift_number(m):
            value = int(m.group(0))
            return str(value + increment if value >= 16 else value)
        return '[' + re.sub(r'\d+', shift_number, inner) + ']'
    return re.sub(r'\[([^\]]+)\]', repl, text)


doc = Document(SOURCE)

# Existing references 16 onward move by two positions to make room for DINOv2
# and UNI at their first appearance in the Background.
reference_heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == 'References')
for p in doc.paragraphs[:reference_heading_index]:
    updated = shift_citations(p.text)
    if updated != p.text:
        replace_paragraph(p, updated)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                updated = shift_citations(p.text)
                if updated != p.text:
                    replace_paragraph(p, updated)

replacements = {
    'Background. Routine haematoxylin and eosin': (
        'Background. Routine haematoxylin and eosin (H&E) slides preserve morphological effects of tumour genotype and the microenvironment. We compared three released whole-slide representation pipelines to identify which driver mutations, fusions, microsatellite-instability measures, genome-doubling and aneuploidy features, oncogenic-pathway states, and derived immune, inflammatory and genomic-context phenotypes could be predicted within each cancer.',
        'Background.'
    ),
    'Methods. We analysed 8,241 patients': (
        'Methods. We analysed 8,241 patients across 32 The Cancer Genome Atlas (TCGA) cancer types and evaluated 1,933 cancer-endpoint pairs with TITAN, Giga-SSL and Prov-GigaPath representations. Multiple slides were mean-pooled within each patient, and identical outcome subsets and nested patient-level folds were used for all three pipelines. A common partial least-squares regression and linear discriminant classification probe selected 1 to 10 components. Cross-validated Q² and area under the receiver operating characteristic curve (AUROC) were the primary metrics; Q² of at least 0.20 and AUROC of at least 0.60 summarized catalogue breadth. Alternative partitions and grouping by TCGA tissue-source-site code assessed internal stability. A separate 9,404-patient TITAN screen used permutation testing and false-discovery-rate control. PathoFMPred was developed as a secondary research interface and registry for fitted models from a defined subset of tasks.',
        'Methods.'
    ),
    'Results. Excluding 11 same-H&E tasks': (
        'Results. After 11 same-H&E tasks were excluded, TITAN, Prov-GigaPath and Giga-SSL reached the Q² threshold in 197, 119 and 94 of 1,496 cross-modal continuous tasks. The corresponding AUROC-threshold counts were 237, 180 and 170 of 426 binary tasks, including 142, 101 and 95 of 243 directly observed genomic-alteration tasks. Signals retained by all three pipelines included THYM GTF2I mutation (AUROC 0.904, 0.893 and 0.884), strict COAD microsatellite instability (0.940, 0.873 and 0.878), UCEC fusion status (0.859, 0.705 and 0.715) and TGCT TGF-beta response (Q² 0.657, 0.623 and 0.642), reported in the order TITAN, Prov-GigaPath and Giga-SSL. The leading pipeline varied by endpoint. In the supporting TITAN screen, grouping by tissue-source-site code moved 83 of 323 candidates below their original threshold and reduced COAD APC and READ APC performance towards chance.',
        'Results.'
    ),
    'Conclusions. Histology contained reproducible': (
        'Conclusions. The three representations captured shared histological signals for mutations, fusions, microsatellite instability, genomic context, sequencing-derived burdens and selected computational immune phenotypes. TITAN covered the largest number of tasks under the common probe, while Giga-SSL or Prov-GigaPath performed best for specific endpoints. This atlas identifies candidates for independent testing. PathoFMPred records the corresponding research models, but the internal TCGA estimates and their sensitivity to cohort structure do not support clinical use.',
        'Conclusions.'
    ),
    'Digital pathology converts routine': (
        'Routine H&E sections record tissue architecture, tumour-cell morphology and the composition of the surrounding stroma. Deep neural networks can learn associations between these patterns and molecular alterations. Coudray et al. first showed that a convolutional neural network could predict selected lung-cancer mutations from histology [1]. Later studies predicted microsatellite instability, including external validation in colorectal cancer [2,3], and extended histology-based prediction to driver mutations and multi-omic features across TCGA cancers [4,5,8,11,13]. Other work inferred RNA expression [6,12], quantified institution-associated bias [7], detected gene fusions [9,10], estimated homologous-recombination deficiency [14] and modelled tumour-microenvironment phenotypes [15]. Together, these studies show that routine histology contains molecular and microenvironmental correlates, although their strength depends on the cancer, outcome and validation cohort.',
        None
    ),
    'Breadth alone is not the gap.': (
        'Foundation models offer a different way to use these images. Instead of training a complete neural network for each endpoint, a foundation model learns a general representation from a large image collection, often without outcome labels. DINOv2 provides a widely used example in computer vision: it trains a vision transformer by self-supervision and produces transferable image features that can support simple downstream models [16]. Pathology foundation models apply this principle to gigapixel slides. They first encode small tissue regions and then aggregate thousands of regional vectors into a slide representation. Histology-specific models such as UNI showed that large-scale self-supervised pretraining can transfer across tissues and diagnostic tasks [17]. A fixed slide representation can therefore be reused for many outcomes, which is valuable when molecularly labelled cohorts are much smaller than the image collections used for pretraining.',
        None
    ),
    'What remains unmapped is how': (
        'Earlier studies already established the breadth of pan-cancer prediction. Fu et al. analysed 17,355 slides across 28 cancers [4], Kather et al. studied more than 5,000 patients across 14 cancers [5], Saldanha et al. externally tested mutation models in seven matched TCGA and CPTAC cancers [11], and Arslan et al. trained 12,093 models for 4,031 genomic, transcriptomic, proteomic and clinical biomarkers in 8,890 TCGA patients across the same 32 cancers [13]. Foundation-model benchmarks have compared many representations, but mainly across smaller sets of classification tasks [18,45,46]. We therefore used one patient-level design to compare three released slide-representation pipelines across a broad catalogue of cancer-specific molecular and derived immune phenotypes. The analysis retains tasks that failed the performance threshold and tasks that lacked sufficient sample size, allowing readers to see both where histological signal was detected and where it was not.',
        None
    ),
    'We evaluated three released representation pipelines': (
        'We compared three released representation pipelines, including their original preprocessing, resolution assumptions, encoder architecture, released layer and pretraining exposure. TITAN combines visual self-supervision with vision-language alignment; its published Mass-340K pretraining corpus excluded TCGA, although TCGA was used for downstream evaluation [18]. Giga-SSL learns a whole-slide representation by self-supervision and its official repository provides 512-dimensional TCGA embeddings [41,43]. Because TCGA contributed to Giga-SSL development, direct exposure to the evaluated images cannot be excluded. Prov-GigaPath combines a DINOv2-pretrained tile encoder with a long-context slide encoder trained on Providence health-system slides; we used the final 768-dimensional slide layer from the public TCGA embedding release [42,44]. We did not fine-tune these encoders or use molecular labels during representation learning. The benchmark therefore compares complete released pipelines under the same 1 to 10 component PLS-based probe, not the intrinsic quality of isolated foundation models.',
        None
    ),
    'In this study, a cancer-endpoint pair': (
        'A cancer-endpoint pair combines one TCGA cancer with one outcome, such as APC mutation in COAD or the wound-healing expression signature in LIHC. Each pair defines a separate cancer-specific prediction task. Patients were included only when the required outcome was available, and a separate model was fitted for each representation. The matched benchmark used the same patients and validation folds for TITAN, Giga-SSL and Prov-GigaPath. Cancers were never pooled into one prediction model.',
        None
    ),
    'Our primary objective was to identify': (
        'Our primary aim was to identify which tumour features could be predicted from routine histology within each cancer and which signals were shared across TITAN, Giga-SSL and Prov-GigaPath. The catalogue included mutations, fusions, microsatellite instability (MSI), genome doubling, aneuploidy, oncogenic pathways, sequencing-derived burdens and derived immune, inflammatory and genomic-context phenotypes. We then examined stability across alternative partitions and grouping by TCGA tissue-source-site code. Throughout the manuscript, predictability means cross-validated agreement with the supplied reference label. It does not imply causality, recovery of the originating assay or clinical interchangeability. PathoFMPred was created as a secondary research tool so that compatible fitted models can be applied and compared without rebuilding the analysis pipeline [34]. The package is important for reuse, but the biological atlas and comparison of the three representation pipelines remain the main contribution.',
        None
    ),
    'Slide vectors of 768 dimensions were supplied': (
        'Slide vectors with 768 dimensions were supplied by TITAN and Prov-GigaPath, and 512-dimensional vectors were supplied by Giga-SSL. TITAN used a CONCH v1.5 tile encoder, its region-of-interest workflow and a released slide-level representation. Giga-SSL used a ResNet-18 tile encoder and sparse-convolutional slide encoder. Prov-GigaPath used a tile encoder pretrained with the DINOv2 self-supervised framework [16] and a long-context slide encoder. Five inclusion criteria were applied: a ready-to-use slide-level TCGA embedding artifact had to be available without reprocessing whole-slide pixels; slide identifiers had to support deterministic linkage to patient and cancer; one fixed-length vector had to be provided per slide; sufficient coverage had to remain for a common-patient analysis across 32 cancers; and reproducible release metadata had to describe a distinct published whole-slide representation strategy. TITAN, Giga-SSL and Prov-GigaPath met all five criteria. Other models that lacked a compatible slide-level TCGA artifact, identifiers or common-cohort coverage are documented in the supplementary inclusion audit. Their exclusion was not interpreted as a judgement of model quality. Upstream preprocessing, spatial resolution, training exposure and released layers were not harmonised. The analysis therefore estimates the performance of each complete released pipeline under the common probe. The reported TITAN Mass-340K and Prov-GigaPath Providence pretraining corpora excluded TCGA. TCGA was used during Giga-SSL development, so downstream molecular labels were held out even though representation learning may have included the evaluated images.',
        None
    ),
    'Outcome labels were acquired from published': (
        'Outcome labels were obtained from published TCGA companion resources and joined by the 12-character participant barcode after representation-specific patient vectors had been constructed. TCGA case and sample metadata were verified through the Genomic Data Commons Cases API [35]. When a source supplied a 15-character TCGA sample barcode, primary-tumour sample type 01 was retained and eligible molecular rows were collapsed to one participant-level outcome. The pan-cancer immune resource published by Thorsson et al. supplied participant-level identifiers only [21]. Source-specific and endpoint-specific missingness was preserved; absence from a molecular source was not interpreted as a negative class. An exact 15-character match identifies the same TCGA primary-tumour specimen, but it does not prove that the slide and molecular assay used the same block, portion, analyte, aliquot, tumour region or subclone. For every outcome group, the source file, worksheet, identifier, filter, aggregation rule, missingness rule and transformation are recorded in the supplementary endpoint dictionary and outcome_source_acquisition_map.csv.',
        None
    ),
    'Immune, inflammatory and genomic-context features were obtained': (
        'Immune, inflammatory and genomic-context phenotypes were imported from the PanImmune_MS worksheet published with the pan-cancer immune analysis by Thorsson et al. [21]. We selected 50 fields that were neither survival outcomes, molecular subtypes nor aggregate labels and used their published participant-level values without recalculation. The set included 22 relative immune-cell fractions inferred from bulk RNA sequencing with CIBERSORT and the LM22 reference [39]; one methylation-derived leukocyte fraction; ten bulk-RNA expression signatures, including IFN-gamma response, TGF-beta response, wound healing, macrophage regulation, lymphocyte infiltration and Th1, Th2 and Th17 programmes; six B-cell and T-cell receptor repertoire metrics; two somatic mutation rates; two in-silico neoantigen burdens; six genomic-context quantities; and the H&E-derived TIL Regional Fraction reported by Saltz et al. [40]. Endpoint-specific missing values were retained. A log1p transformation was applied only to silent and nonsilent mutation rates, SNV and indel neoantigen counts and the number of segments. Every other phenotype from the Thorsson et al. resource was analysed on its published scale. CIBERSORT fractions and expression signatures were interpreted as computational reference phenotypes, not as direct immune-cell measurements.',
        None
    ),
    'The matched benchmark revealed its strongest': (
        'Direct genomic alterations produced the clearest and broadest signal in the matched benchmark. At least one representation reached the descriptive AUROC threshold in 163 of 243 mutation or fusion tasks, and all three did so in 68. For composite genomic-context status, including MSI, genome doubling and oncogenic-pathway states, the corresponding counts were 110 of 183 and 58 of 183. Continuous outcomes were more selective. At least one representation reached Q² of 0.20 in 38 of 413 sequencing-derived burden tasks, 104 of 279 transcriptomic-signature tasks, 29 of 638 computationally inferred immune-cell-fraction tasks and 41 of 166 continuous genomic-context tasks. Table 1 reports the all-three counts for each class. Across these six classes, only 24% to 53% of tasks retained by at least one representation were retained by all three, showing that many signals depended on the representation. The 11 same-H&E TIL Regional Fraction tasks are reported separately. Figure 2A summarizes the biological breadth of the atlas.',
        None
    ),
    'The continuous results linked histology': (
        'Histology predicted selected inflammatory and tissue-context programmes rather than every immune-related phenotype (Figure 2C). All three representations captured the TGF-beta response signature in TGCT at Q² 0.657/0.642/0.623, the Th17 expression programme in THYM at 0.620/0.494/0.500, the methylation-derived leukocyte fraction in BLCA at 0.399/0.395/0.394 and KIRP at 0.368/0.377/0.399, and T-cell receptor diversity in THYM at 0.443/0.429/0.280 for TITAN/Giga-SSL/Prov-GigaPath. Shared signals also included cancer-specific proliferation, macrophage-regulation, lymphocyte-infiltration, Th1, TGF-beta and stromal-fraction phenotypes. These results measure agreement with the computational or sequencing-derived reference values published in the Thorsson et al. pan-cancer immune resource [21]. They do not demonstrate direct recovery of immune-cell abundance. All three representations also reached the threshold in all 11 same-H&E TIL Regional Fraction tasks, which were excluded from the default cross-modal summary. Supplementary Tables S5 to S7 report selected continuous results, feature-family summaries and literature context.',
        None
    ),
    'Previous studies establish that the individual': (
        'The atlas extends several established lines of histology-based prediction. Fu et al. reported whole-genome duplication, chromosomal aneuploidy, driver alterations and tumour-composition signals [4], while Loeffler et al. studied driver genes and oncogenic pathways [8]. Dadhania et al. predicted ERG rearrangement in prostate cancer, and Mayer et al. predicted ALK and ROS1 fusions in lung cancer [9,10]. HE2RNA and later regression studies modelled transcriptomic reference phenotypes [6,12]. HistoTME compared three feature extractors across 30 tumour-microenvironment signatures within one cancer [15], and Arslan et al. screened thousands of multi-omic biomarkers across all 32 TCGA cancers [13]. Foundation-model benchmarks have addressed a different scale of comparison: Neidlinger et al. evaluated 19 models across 31 weakly supervised tasks with external validation in five cohorts [45], Patho-Bench standardized 42 tasks across five slide encoders [46], and the TITAN study compared slide encoders across 62 tasks, including 39 molecular classifications [18].',
        None
    ),
    'Our principal innovation is the scale': (
        'Our study brings these questions into one patient-level comparison of 1,933 cancer-endpoint pairs. The catalogue combines direct genomic alterations with continuous sequencing-derived and computational reference phenotypes, retains tested-negative results and applies the same downstream probe to three released slide representations. Patient-level slide aggregation, matched folds and target-level sensitivity analyses make the representation comparison directly interpretable. The study also shows which biological signals are shared across pipelines and which depend on representation, partition or cohort structure.',
        None
    ),
    'The targeted narrative literature audit': (
        'A targeted narrative cross-check found prior statistical support for 38 of the 41 cancer-gene pairs retained in the TITAN permutation/FDR screen. This recovery supports the biological credibility of the atlas. THYM GTF2I was not identified in the reviewed prediction literature and warrants independent testing, especially because GTF2I mutation is strongly associated with spindle-cell type A and AB thymoma morphology [37]. UCEC and LGG fusion burden and several cancer-specific any-fusion tasks also appear less frequently in published histology-prediction studies, but their novelty and transportability require systematic review and external validation. PathoFMPred provides a practical route for such research testing by applying a compatible fitted model without redistributing patient-level training embeddings. Each result is accompanied by its outcome definition, internal performance and cohort-structure sensitivity metadata.',
        None
    ),
    'Under the common PLS-based probe': (
        'Under the common PLS-based probe, TITAN, Giga-SSL and Prov-GigaPath captured cancer-specific histological associations with mutations, fusions, MSI, genome doubling, aneuploidy, oncogenic pathways, sequencing-derived burdens and selected derived immune, inflammatory and genomic-context phenotypes. At least one representation reached the descriptive threshold in 496 cancer-endpoint pairs covering 144 feature definitions, and all three did so in 213 pairs. Strong shared priorities included THYM GTF2I, THCA BRAF, LGG IDH1 and TP53, BLCA FGFR3, strict COAD MSI, the TGCT TGF-beta response signature, the THYM Th17 programme and T-cell receptor diversity, and leukocyte fraction in BLCA and KIRP. TITAN produced the highest observed effect in 345 of 485 cross-modal threshold-reaching pairs, while Giga-SSL and Prov-GigaPath led meaningful endpoint-specific subsets. The atlas therefore supports tumour-feature-specific representation choice and independent biological validation. PathoFMPred provides a secondary research interface for applying compatible models.',
        None
    ),
}

for p in doc.paragraphs:
    for start, (text, bold_prefix) in replacements.items():
        if p.text.startswith(start):
            replace_paragraph(p, text, bold_prefix)
            break

# Preserve author details requested in the current project record and cite the
# principal statistical methods where they are first described.
for p in doc.paragraphs:
    if p.text.startswith('4 Department of Surgery, Faculty of Health Sciences'):
        replace_paragraph(
            p,
            '4 Department of Surgery, School of Clinical Medicine, Faculty of Health Sciences, University of the Witwatersrand, Johannesburg, Gauteng, South Africa'
        )
    elif p.text.startswith('Each eligible cancer-endpoint pair was analysed'):
        text = p.text.replace(
            'For continuous outcomes, PLS regression was fitted,',
            'For continuous outcomes, PLS regression was fitted [19],'
        ).replace(
            'For binary outcomes, PLS-LDA was fitted,',
            'For binary outcomes, PLS-LDA was fitted [20],'
        )
        replace_paragraph(p, text)
    elif p.text.startswith('A total of 2,073 eligible pairs were evaluated'):
        replace_paragraph(
            p,
            p.text.replace('Benjamini-Hochberg correction was applied', 'Benjamini-Hochberg correction [29] was applied')
        )
    elif p.text.startswith('The analyses were performed in R 4.6.0'):
        replace_paragraph(p, p.text + ' Reporting was audited against TRIPOD+AI [30].')

# Figure 2 caption now describes the redesigned grouped bars in panel A.
for p in doc.paragraphs:
    if p.text.startswith('Figure 2. Biological features predictable'):
        replace_paragraph(
            p,
            'Figure 2. Biological features predicted from the three released pathology representation pipelines. Panel A shows, for each tumour-feature class, the percentage and number of eligible cancer-endpoint pairs that reached Q² of at least 0.20 or AUROC of at least 0.60 in at least one representation and in all three representations. The 11 same-H&E TIL-fraction tasks are excluded. Panel B presents selected direct genomic and genomic-context examples, and panel C presents selected immune, inflammatory and tissue-context reference phenotypes. Every displayed example reached the threshold in all three representations and retained it after grouping by TCGA tissue-source-site code. Points and numerical values follow the order TITAN, Giga-SSL and Prov-GigaPath. The coloured label identifies the pipeline with the highest observed effect for that cancer-endpoint pair. Grouped retention is an internal cohort-structure sensitivity, not external validation.'
        )

# Clarify the source name in the linkage table.
for table in doc.tables:
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == 'Thorsson PanImmune':
            replace_paragraph(row.cells[0].paragraphs[0], 'Pan-cancer immune resource (Thorsson et al.)')

# Remove accidental doubled punctuation and missing spaces introduced in the edited source.
for p in doc.paragraphs:
    text = p.text
    fixed = text.replace('..', '.').replace('TCGA.Because', 'TCGA. Because')
    fixed = fixed.replace('immune feature(Figure', 'immune feature (Figure')
    fixed = fixed.replace('tasks.These', 'tasks. These')
    fixed = fixed.replace('GigaPath  AUROCs', 'GigaPath AUROCs')
    if fixed.startswith('. Across'):
        fixed = fixed[2:]
    if fixed != text:
        replace_paragraph(p, fixed)

# Rebuild the numbered bibliography with DINOv2 and UNI inserted at positions 16 and 17.
paragraphs = doc.paragraphs
ref_heading_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == 'References')
old_ref_paragraphs = paragraphs[ref_heading_index + 1:]
old_refs = [p.text for p in old_ref_paragraphs if re.match(r'^\d+\.\s', p.text)]
assert len(old_refs) == 44, len(old_refs)
new_refs = old_refs[:15] + [
    '16. Oquab M, et al. DINOv2: learning robust visual features without supervision. Trans Mach Learn Res. 2024. https://openreview.net/forum?id=a68SUt6zFt.',
    '17. Chen RJ, et al. Towards a general-purpose foundation model for computational pathology. Nat Med. 2024;30:850–862. doi:10.1038/s41591-024-02857-3.'
]
for item in old_refs[15:]:
    number, rest = item.split('.', 1)
    new_refs.append(f'{int(number) + 2}.{rest}')

for p, text in zip(old_ref_paragraphs, new_refs[:len(old_ref_paragraphs)]):
    replace_paragraph(p, text)
last = old_ref_paragraphs[-1]
for text in new_refs[len(old_ref_paragraphs):]:
    last = insert_after(last, text)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

# Replace the Figure 2 raster without changing its relationship or inline size.
with tempfile.TemporaryDirectory() as td:
    tmp = Path(td) / 'repacked.docx'
    with zipfile.ZipFile(OUTPUT, 'r') as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = FIGURE2.read_bytes() if item.filename == 'word/media/image2.png' else zin.read(item.filename)
            zout.writestr(item, data)
    shutil.move(tmp, OUTPUT)

print(OUTPUT)
