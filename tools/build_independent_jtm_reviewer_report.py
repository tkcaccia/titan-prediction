#!/usr/bin/env python3
"""Build a fresh Journal of Translational Medicine reviewer report."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "manuscript" / "independent_reviewer_report_JTM.docx"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
INK = "202124"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C3D0", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_fixed_table_layout(table, widths_inches):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total_dxa = int(round(sum(widths_inches) * 1440))
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = Inches(widths_inches[idx])
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(int(round(widths_inches[idx] * 1440))))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, size=None, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_body(doc, text, bold_lead=None, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True, color=INK)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11, italic=italic, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, italic=italic, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_major_comment(doc, number, title, paragraphs, actions=None):
    add_heading(doc, f"{number}. {title}", level=2)
    for paragraph in paragraphs:
        add_body(doc, paragraph)
    if actions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.10
        lead = p.add_run("Required revision. ")
        set_run_font(lead, size=11, bold=True, color=RED)
        rest = p.add_run(actions)
        set_run_font(rest, size=11, color=INK)


def add_minor_comment(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    r1 = p.add_run(f"{number}. ")
    set_run_font(r1, size=10.5, bold=True, color=BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.80)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, NAVY, 14, 7),
        "Heading 2": (13, BLUE, 11, 5),
        "Heading 3": (11.5, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Journal of Translational Medicine | Confidential peer review")
    set_run_font(hr, size=8.5, color=MID_GRAY)
    footer = section.footer
    add_page_number(footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(5)
    kr = kicker.add_run("JOURNAL-STYLE PEER REVIEW")
    set_run_font(kr, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    tr = title.add_run("Independent Reviewer Report")
    set_run_font(tr, size=24, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run(
        "Patient-level comparison of three released pathology representation pipelines with a common PLS-based probe across 32 TCGA cancers"
    )
    set_run_font(sr, size=12.5, italic=True, color=MID_GRAY)

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_fixed_table_layout(meta, [1.55, 5.19])
    set_table_borders(meta, color="D5DDE5", size="5")
    rows = [
        ("Journal", "Journal of Translational Medicine"),
        ("Article type", "Research Article, Molecular Pathology"),
        ("Recommendation", "Major revision; not suitable for acceptance in its current form"),
        ("Material reviewed", "Main manuscript and Supplementary Information"),
        ("Review scope", "Scientific, statistical, translational, reproducibility and presentation review; the full R pipeline and private fitted objects were not independently rerun"),
    ]
    for idx, (label, value) in enumerate(rows):
        c0, c1 = meta.rows[idx].cells
        c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c0, LIGHT_BLUE)
        for p in c0.paragraphs + c1.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        r0 = c0.paragraphs[0].add_run(label)
        set_run_font(r0, size=9.5, bold=True, color=NAVY)
        r1 = c1.paragraphs[0].add_run(value)
        set_run_font(r1, size=9.5, color=INK)

    add_heading(doc, "Overall assessment", level=1)
    add_body(
        doc,
        "This manuscript presents an ambitious patient-level benchmark of three released pathology representation pipelines, TITAN, Giga-SSL and Prov-GigaPath, across 8,241 shared TCGA patients and 1,933 cancer-endpoint prediction tasks. The authors compare the representations on identical patients and validation folds with a common PLS-based downstream probe. They also provide a deeper TITAN-only screen with permutation and false-discovery-rate filtering, extensive internal sensitivity analyses and a software and model-registry layer."
    )
    add_body(
        doc,
        "The work has important strengths. It uses the patient as the unit of analysis and cross-validation, pools multiple slides before outcome matching, preserves missing and ineligible outcomes, audits exact slide-set overlap, distinguishes directly observed genomic alterations from derived reference phenotypes, and tests sensitivity to alternative partitions and grouping by TCGA tissue-source-site code. The manuscript is also unusually candid about the descriptive nature of the matched benchmark, the absence of external validation, the limits of slide and molecular linkage, the uncalibrated binary scores and the restricted availability of fitted objects."
    )
    add_body(
        doc,
        "The topic fits the journal's Medical Bioinformatics, Molecular Pathology, Disease Biomarkers and Translational Imaging scope. However, the current evidence remains internal to TCGA, several headline results depend on the downstream probe and pragmatic thresholds, many grouped analyses have sparse or single-class folds, pathology quality was not independently adjudicated, and most continuous findings concern computational reference phenotypes rather than direct measurements. The fitted-model resource is also incomplete and lacks a durable reader-facing access route. I therefore recommend major revision."
    )

    add_heading(doc, "Principal strengths", level=1)
    strengths = [
        ("Patient-level design.", "Slides from one participant are aggregated before outcome linkage and remain within one validation fold, which reduces the leakage risk created by repeated slides."),
        ("Matched representation comparison.", "The three pipelines use the same 8,241-patient cohort, target-labelled subsets, folds and seeds, and the exact-common-slide analysis shows that small slide-set differences do not explain the aggregate ranking."),
        ("Coherent binary analysis.", "The primary matched benchmark tunes components with inner out-of-fold AUROC and evaluates outer out-of-fold AUROC, while deriving operating-point metrics from a training-only threshold."),
        ("Outcome provenance.", "The authors separate direct genomic alterations, sequencing-derived burdens, transcriptomic signatures, inferred immune fractions, composite scores and the same-H&E TIL fraction."),
        ("Internal robustness auditing.", "Alternative partitions, tissue-source-site-code grouping, matched-random controls, component-range sensitivity, class-size audits and pooling sensitivities expose several important limitations rather than concealing them."),
        ("Transparency.", "The Supplement records source-specific missingness, model eligibility, component selection, permutation status, grouped-fold composition, software metadata and model-access restrictions."),
    ]
    for label, detail in strengths:
        add_body(doc, f"{label} {detail}", bold_lead=label)

    add_heading(doc, "Major comments", level=1)

    add_major_comment(
        doc,
        1,
        "Independent external validation is the central missing element",
        [
            "All performance estimates come from TCGA. Alternative partitions, exact-slide intersections and tissue-source-site-code grouping are useful internal checks, but they cannot establish transportability to another health system, scanner mix, staining distribution, population or molecular laboratory. The prospectively locked UCEC protocol in Supplementary Table S6c is good planning, but it contains no external result and therefore does not address the principal translational limitation.",
            "This point is decisive for Journal of Translational Medicine because the manuscript presents a research prioritisation resource that is intended to guide which molecular associations warrant further testing. At present, the paper establishes internal association, not external prediction or clinical utility.",
        ],
        "Evaluate a small locked and clinically interpretable set in an independent cohort using the exact released representation pipeline and model object, without refitting, recalibration or threshold adjustment. The set should span at least one common mutation, one MSI or genomic-instability endpoint and one continuous sequence-derived or transcriptomic endpoint. Report every prespecified target, including failures and non-evaluable outcomes. If external validation is not feasible, restrict the title, abstract, software language and conclusions to an internal TCGA benchmark and consider whether a computational pathology or bioinformatics resource venue is more appropriate."
    )

    add_major_comment(
        doc,
        2,
        "The study compares released pipeline and probe combinations, not intrinsic foundation-model quality",
        [
            "The manuscript appropriately states that the representations differ in preprocessing, spatial resolution, tile and slide encoders, released layer, dimensionality and pretraining exposure. Giga-SSL development used TCGA, and direct overlap between representation-learning images and evaluated slides cannot be excluded. TITAN and Prov-GigaPath were reported to exclude TCGA from pretraining. The downstream labels are held out for all three models, but the image-exposure conditions are not equivalent.",
            "The downstream method also affects the apparent representation leader. In the fixed 47-target ridge sensitivity, ridge retained the PLS-defined leading representation for only 22 of 35 binary and 8 of 12 continuous targets. Expanding the component grid changed the leading representation for 43 of 366 selected binary pairs. These findings are central because the title and primary result compare representations.",
        ],
        "Keep every comparative claim explicitly conditional on the released pipeline and the specified PLS-based probe. Add a broader symmetric ridge sensitivity or another fixed linear baseline using identical folds and objectives, preferably across the complete matched atlas or at least every union-crossing and near-threshold task. Extend the component-range audit to the remaining binary tasks and an adequately defined continuous subset. The main comparison table should display pretraining exposure, dimensionality, component-ceiling frequency and numerical-failure policy beside the performance summaries."
    )

    add_major_comment(
        doc,
        3,
        "The matched atlas is descriptive and threshold dependent",
        [
            "The primary three-representation benchmark reports effect-threshold crossings at Q² at least 0.20 or AUROC at least 0.60. It does not apply representation-specific permutation testing or multiplicity correction. The reported 496 union crossings and 213 all-three crossings are therefore catalogue-navigation counts, not statistically screened discoveries. This distinction is stated, but the abstract and Results still devote considerable emphasis to the counts.",
            "The alternative-partition analysis also shows that hard support categories are unstable. Exactly-two and representation-specific classes are much less reproducible than all-three consensus, and the selected task set contains crossing or near-threshold pairs rather than the complete atlas. A change of a few hundredths around an arbitrary threshold can alter the category without materially changing the effect.",
        ],
        "Make paired Q² and AUROC values, alternative-partition variability, rank stability and grouped-minus-matched-random differences the primary evidence. Use crossing labels only as database tags. Report repeat distributions and crossing proportions beside every highlighted result, including continuous examples. If the authors wish to make inferential claims about representation-specific associations, perform a clearly defined permutation and multiplicity analysis on a restricted target set."
    )

    add_major_comment(
        doc,
        4,
        "Tissue-source-site-code sensitivity is important, but many grouped folds are inadequate",
        [
            "The grouped analysis is one of the strongest parts of the study. It shows marked attenuation for some important results, including READ-APC, and the matched-random controls help separate the consequences of fold size and class balance from additional cohort-structure shifts. The interpretation also correctly avoids treating the barcode-derived code as a scanner, laboratory or institution identifier.",
            "However, the Supplement reports sparse-fold warnings for 162 of 273 binary tasks and 51 of 223 continuous tasks. Twenty-eight binary tasks contain a single-class outer test fold, 120 contain a single-class inner validation fold, and the minimum inner training count reaches zero for one class. A pooled outer metric can still be calculated, but retention labels from these designs do not have uniform reliability.",
        ],
        "Report grouped-retention summaries separately for tasks that pass and fail the fold-adequacy audit. For each highlighted task, provide contributing code count, realised fold count, minimum training-class counts and any single-class fold. Do not assign an unqualified robustness interpretation when the grouped analysis is sparse. Keep the matched-random comparator beside every highlighted grouped estimate."
    )

    add_major_comment(
        doc,
        5,
        "Small binary classes contribute substantially to the headline counts",
        [
            "The inclusive eligibility rule requires only 20 patients per binary class. In the matched atlas, the larger-sample stratum accounts for 147 of 237 TITAN binary crossings, 108 of 170 Giga-SSL crossings and 117 of 180 Prov-GigaPath crossings. The remaining 90, 62 and 63 crossings arise from the smaller sample-size stratum. This is a large fraction of the headline binary breadth, not a peripheral issue.",
            "The manuscript appropriately excludes smaller-class models from default inference and reports learning curves and component stability for the TITAN layer. Nevertheless, a reader can easily interpret the inclusive counts as equally mature associations. Binary output is also an uncalibrated score rather than a probability, and PPV and NPV are conditional on TCGA prevalence.",
        ],
        "Present inclusive and larger-sample-stratum counts in parallel in the abstract, main table and key figures. Restrict biological prioritisation and software defaults to adequately sized tasks unless an exploratory label is visually unavoidable. For highlighted binary endpoints, report prevalence, PR-AUC with its no-skill reference, sensitivity, specificity, PPV, NPV and the exact training-derived operating rule. Do not describe rank or score as probability."
    )

    add_major_comment(
        doc,
        6,
        "Pathology quality and molecular-slide linkage remain insufficient for translational interpretation",
        [
            "The patient-first pooling strategy is sound, and the mean-versus-median, first-slide, leave-one-slide-out and 30-slide SARC sensitivities are informative. They do not establish that the pooled slides contain representative tumour. The source data lack structured tumour cellularity, tissue area, artefact, biopsy or resection status and slide-quality fields. The generated-text audit flagged six patients with no-residual-tumour language, including one 30-slide SARC patient, but those narratives are not independent pathology annotations.",
            "Molecular linkage also remains imperfect. Thorsson phenotypes are linked only at participant level. A 15-character sample-barcode match for other sources does not prove that the slide and assay used the same block, portion, aliquot, tumour region or subclone. These limitations affect every cross-modal association and are especially important for heterogeneous tumours and multi-slide patients.",
        ],
        "Obtain blinded pathologist quality review for a manageable set of highlighted slides and report tumour presence, approximate tumour area, major artefact and specimen type. At minimum, verify the slides supporting the principal direct genomic and MSI examples and repeat those analyses after excluding non-representative material. Retain the generated narrative audit only as a non-adjudicated sensitivity and do not use it as a surrogate pathology label."
    )

    add_major_comment(
        doc,
        7,
        "Derived reference phenotypes dominate the continuous catalogue",
        [
            "The provenance audit is a major strength, but 93% to 97% of continuous crossings arise from the Thorsson-derived catalogue. These outcomes include RNA signatures, CIBERSORT fractions, methylation-derived leukocyte estimates, repertoire measures, in-silico neoantigens, composite genomic-context scores and an H&E-derived TIL fraction. Agreement with these references is not equivalent to direct immune-cell measurement, reconstruction of the originating assay or prediction of a certified biomarker.",
            "The paper now separates the same-H&E TIL fraction and reports same-H&E-excluded rates, which is appropriate. The biological interpretation nevertheless remains limited because the global embeddings do not identify the morphology that supports a prediction, and the nearest-neighbour audit uses generated reports rather than independent pathology review.",
        ],
        "Lead the translational interpretation with direct genomic alterations and sequence-derived outcomes. Present transcriptomic and inferred immune results as agreement with computational reference phenotypes. For a small set of strong cross-representation and tissue-source-site-code-stable targets, add spatial relevance analysis or blinded pathologist review if tile-level data can be regenerated. If this is not possible, state clearly that the study prioritises associations without localising their histological basis."
    )

    add_major_comment(
        doc,
        8,
        "PathoFMPred is not yet a reusable fitted-model resource",
        [
            "The manuscript is commendably explicit that the intended public package shell will contain code, registry, schemas, tests and synthetic fixtures but cannot perform inference without a separately supplied fitted object. The 917 controlled objects cover only part of the matched atlas, use a TITAN-qualified target universe for the alternative representations, and include some representation-specific fits that did not cross the corresponding effect threshold. No durable reader-facing controlled-access route is promised.",
            "The fitted objects also have representation-specific licensing constraints. In particular, the Supplement states that TITAN-derived fitted objects require written clarification or permission before redistribution. These conditions are incompatible with an unrestricted claim of a reusable prediction package.",
        ],
        "Before acceptance, provide reviewers with access sufficient to inspect the package and representative fitted objects. Resolve the licence and redistribution status for each representation, define a durable access route, and state exactly which atlas tasks have fitted objects. If fitted objects cannot be distributed or accessed, describe PathoFMPred as an analysis interface and model registry rather than a functioning public prediction resource. Replace all release, commit and DOI placeholders with an immutable archived version that matches the reported results."
    )

    add_major_comment(
        doc,
        9,
        "The manuscript needs a sharper translational synthesis",
        [
            "The Discussion is careful and methodologically sophisticated, but it remains primarily a benchmark narrative. A translational reader needs a concise answer to which associations should be tested next, why those associations are clinically or biologically relevant, and which findings should be deprioritised because they depend on representation, partition or cohort structure.",
            "The strongest candidates appear to include THYM-GTF2I, THCA-BRAF, LGG-IDH1 or TP53, strict COAD MSI and selected genomic-instability or inflammatory reference phenotypes. However, these examples currently sit among many catalogue summaries and caveats. The paper would benefit from a small, transparent priority table based on effect size, partition stability, sample size, grouped-fold adequacy and endpoint provenance, without turning those dimensions into a clinical grade.",
        ],
        "Add a concise translational-priority table and organize the Discussion around a small number of direct, derived and cohort-sensitive examples. Explain the plausible research use for each target and the next validation step. State once, clearly, that the atlas supports research prioritisation rather than assay replacement or clinical decision making."
    )

    add_heading(doc, "Minor and editorial comments", level=1)
    minor = [
        "The abstract is overloaded with counts and methodological qualifications. Retain the most informative provenance-stratified rates, the common cohort size, the main paired metric, the grouped-sensitivity result and the absence of external validation. Move the remaining counts to the Results.",
        "Figure 1 contains a visible collision at the lower centre: the Prov-GigaPath label, slide count and exact-common-slide note overlap. Rebuild the figure with more bottom margin and separate the note from the circle labels.",
        "Figure 3 and main Table 3 are dense at journal page width. Increase axis and legend text, simplify the panels, and consider moving detailed grouped-fold fields to the Supplement while keeping the key comparison in the main text.",
        "The Supplement contains a blank page 48. Remove the blank page and audit all section breaks after conversion to PDF.",
        "The 82-page Supplement remains difficult to navigate. Several wide tables use very small text, and the machine-readable file inventory occupies the final four pages. Keep interpretive summaries in the PDF and move exhaustive inventories to a repository README or data dictionary.",
        "Supplementary Table S1a is transparent but lengthy. A short statement in the Supplement plus the machine-readable chronology file would be sufficient. The revision history should not compete with the scientific narrative.",
        "Ensure that every reference is cited at the specific statement it supports. The opening Supplementary Methods paragraph currently cites broad reference ranges, which is not a substitute for claim-level citation. Remove references that remain unused after this audit.",
        "Main Table 3 should include prevalence and PR-AUC, or point directly to a compact table that supplies them, because several mutation and fusion tasks are imbalanced.",
        "Keep the distinction between the AUROC-centred matched benchmark and the balanced-accuracy-based supporting TITAN screen visible wherever results from both layers appear on the same page.",
        "Additional files 2 and 3 are described in the main manuscript as TITAN-input reports, whereas Supplementary Figure S7 compares all three representation inputs. Clarify exactly what each PDF contains and ensure that the submitted files match the description.",
        "Complete the competing-interests, funding and author-contribution declarations. Instructions to the authors such as 'must insert' and 'pending confirmation' should not remain in the submitted manuscript.",
        "Insert the final repository release tag, immutable commit and persistent DOI before submission, and verify that the deposited files were generated from that exact snapshot.",
        "The post hoc COAD radar profiles are acceptable as a software illustration in the Supplement, provided that the original units remain visible and every percentile or binary rank is labelled immediately as a reference rank, not a probability.",
        "The title accurately limits the central comparison to three released pipelines under a common PLS-based probe. Retain this restraint in the abstract and conclusions.",
    ]
    for idx, text in enumerate(minor, 1):
        add_minor_comment(doc, idx, text)

    doc.add_page_break()
    add_heading(doc, "Confidential comments to the editor", level=1)
    add_body(
        doc,
        "This is a technically serious and unusually transparent TCGA benchmark. The patient-level pooling and validation design, common-patient three-representation comparison, endpoint-provenance audit, alternative-partition analysis, tissue-source-site-code controls and reporting of negative or ineligible tasks are clear strengths. The subject is compatible with the journal's Medical Bioinformatics and Molecular Pathology remit."
    )
    add_body(
        doc,
        "My concern is translational maturity rather than basic technical competence. No model has been evaluated in an independent cohort. A substantial fraction of grouped analyses have sparse or single-class folds, pathology quality was not independently adjudicated, and most continuous results concern derived computational phenotypes. The fitted-model layer is private, partial and subject to unresolved access and redistribution constraints. The submitted manuscript also contains incomplete declarations and release placeholders."
    )
    add_body(
        doc,
        "I would support major revision if the authors can add a locked external validation and resolve the software and licensing route. If external validation cannot be added, the work may still be valuable as a computational pathology benchmark and research resource, but its fit for a translational medicine journal becomes substantially weaker. I did not independently rerun the full R pipeline or inspect the private fitted objects."
    )

    add_heading(doc, "Summary judgement", level=1)
    summary = doc.add_table(rows=1, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_fixed_table_layout(summary, [2.25, 4.49])
    set_table_borders(summary, color="B7C3D0", size="6")
    for cell, text in zip(summary.rows[0].cells, ["Criterion", "Assessment"]):
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=10, bold=True, color=WHITE)
    rows = [
        ("Scientific question", "Important and timely"),
        ("Internal design", "Strong, patient-level and unusually well audited"),
        ("Representation comparison", "Informative but conditional on pipeline, training exposure and downstream probe"),
        ("Statistical qualification", "Descriptive for the matched atlas; stronger but TITAN-only for permutation and FDR"),
        ("Biological interpretation", "Plausible for selected direct alterations; limited for many derived phenotypes"),
        ("External validation", "Absent"),
        ("Reproducibility", "Strong analysis documentation; fitted-model access remains unresolved"),
        ("Presentation", "Main text generally clear; Supplement overlong with visible layout defects"),
        ("Current decision", "Major revision"),
    ]
    for idx, (criterion, assessment) in enumerate(rows, 1):
        cells = summary.add_row().cells
        if idx % 2 == 0:
            set_cell_shading(cells[0], LIGHT_GRAY)
            set_cell_shading(cells[1], LIGHT_GRAY)
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
        r0 = cells[0].paragraphs[0].add_run(criterion)
        set_run_font(r0, size=9.5, bold=True, color=NAVY)
        r1 = cells[1].paragraphs[0].add_run(assessment)
        set_run_font(r1, size=9.5, color=INK)

    doc.core_properties.title = "Independent reviewer report for Journal of Translational Medicine"
    doc.core_properties.subject = "Fresh peer review of manuscript and Supplementary Information"
    doc.core_properties.author = "Independent reviewer"
    doc.core_properties.keywords = "peer review, Journal of Translational Medicine, computational pathology"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
