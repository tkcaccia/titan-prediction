from pathlib import Path
import shutil

from docx import Document


MANUSCRIPT = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/manuscript/"
    "manuscript_JTM_multifoundation_atlas.docx"
)
BACKUP = MANUSCRIPT.with_name(
    "manuscript_JTM_multifoundation_atlas.before_pathofmpred_intro.docx"
)

OLD_START = "The narrower unresolved gap is a reproducible multi-representation atlas"
NEW_START = "The gap addressed here is a reproducible atlas"
NEW_TEXT = (
    "The gap addressed here is a reproducible atlas comparing molecular and derived immune-feature "
    "associations across released pipelines for the same patients and endpoints. The atlas pools "
    "slides before outcome matching, validates at patient level and preserves negative and ineligible "
    "results. As a secondary output, we developed PathoFMPred [22], an R package for "
    "representation-specific feature validation and permitted fitted linear models. The main "
    "contribution remains the three-pipeline patient-level benchmark of outcome provenance, partition "
    "stability and cohort-structure sensitivity."
)


def replace_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def main():
    if not BACKUP.exists():
        shutil.copy2(MANUSCRIPT, BACKUP)
    document = Document(MANUSCRIPT)
    matches = [
        p for p in document.paragraphs
        if p.text.startswith((OLD_START, NEW_START))
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one Introduction paragraph, found {len(matches)}")
    replace_paragraph_text(matches[0], NEW_TEXT)
    document.save(MANUSCRIPT)

    check = Document(MANUSCRIPT)
    revised = [p.text for p in check.paragraphs if p.text == NEW_TEXT]
    if revised != [NEW_TEXT]:
        raise RuntimeError("The revised Introduction paragraph did not round-trip exactly")
    print(f"Updated: {MANUSCRIPT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
