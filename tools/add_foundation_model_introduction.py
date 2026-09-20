from pathlib import Path
import re
import shutil

from docx import Document


MANUSCRIPT = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/manuscript/"
    "manuscript_JTM_multifoundation_atlas.docx"
)
BACKUP = MANUSCRIPT.with_name(
    "manuscript_JTM_multifoundation_atlas.before_foundation_model_intro.docx"
)

GAP_START = "The gap addressed here is a reproducible multi-representation atlas"
FOUNDATION_START = "Pathology foundation models are neural networks"
FOUNDATION_TEXT = (
    "Pathology foundation models are neural networks pretrained on large histopathology collections "
    "to produce reusable visual representations. A whole-slide pipeline encodes image tiles, "
    "aggregates them into a slide embedding and links this embedding to a smaller endpoint-specific "
    "model, rather than training a new image encoder for each biomarker. Transfer nevertheless depends "
    "on the pretraining data, image resolution, tiling, architecture and slide aggregation [17-21]. "
    "We therefore compare released embedding pipelines under a common probe, not intrinsic "
    "foundation-model quality."
)

# The foundation-model sources move ahead of the PathoFMPred package citation.
RENUMBER = {17: 22, 18: 17, 19: 18, 20: 19, 21: 20, 22: 21}
REFERENCE_ORDER = [18, 19, 20, 21, 22, 17]


def replace_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def remap_bracket_citations(text):
    def remap_group(match):
        content = match.group(1)

        def remap_number(number_match):
            number = int(number_match.group(0))
            return str(RENUMBER.get(number, number))

        return "[" + re.sub(r"\d+", remap_number, content) + "]"

    return re.sub(r"\[([0-9,;\-–\s]+)\]", remap_group, text)


def main():
    shutil.copy2(MANUSCRIPT, BACKUP)
    document = Document(MANUSCRIPT)
    paragraphs = document.paragraphs
    reference_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() == "References"
    )

    # Renumber existing body citations before inserting the new [17-21] citation.
    for paragraph in paragraphs[:reference_index]:
        revised = remap_bracket_citations(paragraph.text)
        if revised != paragraph.text:
            replace_paragraph_text(paragraph, revised)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    revised = remap_bracket_citations(paragraph.text)
                    if revised != paragraph.text:
                        replace_paragraph_text(paragraph, revised)

    gap_matches = [p for p in document.paragraphs if p.text.startswith(GAP_START)]
    if len(gap_matches) != 1:
        raise RuntimeError(f"Expected one study-gap paragraph, found {len(gap_matches)}")
    if not any(p.text.startswith(FOUNDATION_START) for p in document.paragraphs):
        gap_matches[0].insert_paragraph_before(FOUNDATION_TEXT, style=gap_matches[0].style)

    # Reorder bibliography entries 17-22 and update their displayed numbers.
    paragraphs = document.paragraphs
    reference_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() == "References"
    )
    reference_paragraphs = {}
    for paragraph in paragraphs[reference_index + 1:]:
        match = re.match(r"^(\d+)\.\s+(.*)$", paragraph.text)
        if match:
            reference_paragraphs[int(match.group(1))] = match.group(2)

    if not all(number in reference_paragraphs for number in range(17, 23)):
        raise RuntimeError("Could not locate bibliography entries 17-22")

    target_paragraphs = [
        next(
            paragraph for paragraph in paragraphs[reference_index + 1:]
            if re.match(rf"^{number}\.\s", paragraph.text)
        )
        for number in range(17, 23)
    ]
    for new_number, old_number, paragraph in zip(
        range(17, 23), REFERENCE_ORDER, target_paragraphs
    ):
        replace_paragraph_text(
            paragraph,
            f"{new_number}. {reference_paragraphs[old_number]}",
        )

    document.save(MANUSCRIPT)
    print(f"Updated: {MANUSCRIPT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
