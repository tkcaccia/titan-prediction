from pathlib import Path
import shutil

from docx import Document


MANUSCRIPT = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/manuscript/"
    "manuscript_JTM_multifoundation_atlas.docx"
)
BACKUP = MANUSCRIPT.with_name(
    "manuscript_JTM_multifoundation_atlas.before_foundation_gap_condense.docx"
)

FOUNDATION_START = "Pathology foundation models are neural networks"
GAP_START = "The gap addressed here is a reproducible atlas"

FOUNDATION_TEXT = (
    "Pathology foundation models are neural networks pretrained on large histopathology collections "
    "to produce reusable visual representations. A whole-slide pipeline encodes image tiles, "
    "aggregates them into a slide embedding and links this embedding to a smaller endpoint-specific "
    "model, rather than training a new image encoder for each biomarker. Transfer nevertheless depends "
    "on the pretraining data, image resolution, tiling, architecture and slide aggregation [17-21]. "
    "We therefore compare released embedding pipelines under a common probe, not intrinsic "
    "foundation-model quality."
)

GAP_TEXT = (
    "The gap addressed here is a reproducible atlas comparing molecular and derived immune-feature "
    "associations across released pipelines for the same patients and endpoints. The atlas pools "
    "slides before outcome matching, validates at patient level and preserves negative and ineligible "
    "results. As a secondary output, we developed PathoFMPred [22], an R package for "
    "representation-specific feature validation and permitted fitted linear models. The main "
    "contribution remains the three-pipeline patient-level benchmark of outcome provenance, partition "
    "stability and cohort-structure sensitivity."
)


def replace_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def replace_one(document, start, text):
    matches = [p for p in document.paragraphs if p.text.startswith(start)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {start!r}, found {len(matches)}")
    replace_paragraph_text(matches[0], text)


def main():
    shutil.copy2(MANUSCRIPT, BACKUP)
    document = Document(MANUSCRIPT)
    replace_one(document, FOUNDATION_START, FOUNDATION_TEXT)
    replace_one(document, GAP_START, GAP_TEXT)
    document.save(MANUSCRIPT)
    print(f"Updated: {MANUSCRIPT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
