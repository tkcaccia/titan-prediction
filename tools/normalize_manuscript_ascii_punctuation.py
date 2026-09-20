#!/usr/bin/env python3
"""Normalize manuscript punctuation while preserving DOCX formatting."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document


EXACT_REPLACEMENTS = {
    "PLS1-versus-PLS2": "endpoint-wise versus multi-outcome PLS",
    "PLS1": "endpoint-wise PLS",
    "PLS2": "multi-outcome PLS",
    "Research article — Molecular Pathology | Journal of Translational Medicine":
        "Research article: Molecular Pathology | Journal of Translational Medicine",
    "Here ‘predictability’ means cross-validated agreement with the supplied reference label—not causality, mechanism, analytical recovery of the originating assay or assay replacement.":
        "Here 'predictability' means cross-validated agreement with the supplied reference label. It does not imply causality, mechanism, analytical recovery of the originating assay or assay replacement.",
    "the patient—not the slide—was the unit of analysis and cross-validation":
        "the patient, rather than the slide, was the unit of analysis and cross-validation",
    "Excluding them—the default continuous summary—crossings were":
        "In the default continuous summary, which excluded them, crossings were",
    "explicit negative and sensitivity results—not unprecedented mutation targets or a universal foundation-model ranking":
        "explicit negative and sensitivity results. It does not claim unprecedented mutation targets or a universal foundation-model ranking",
}


ASCII_PUNCTUATION = str.maketrans({
    "–": "-",
    "‑": "-",
    "−": "-",
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "…": "...",
})


def iter_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def normalize_run_text(text: str) -> str:
    for old, new in EXACT_REPLACEMENTS.items():
        text = text.replace(old, new)
    # Convert any remaining em-dash construction to an ordinary comma phrase.
    # This catches compound wording that is not practical to enumerate above.
    text = re.sub(r"\s*—\s*", ", ", text)
    text = text.translate(ASCII_PUNCTUATION)
    return text


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()

    if args.backup:
        shutil.copy2(args.input, args.backup)

    document = Document(args.input)
    changed_runs = 0
    for paragraph in iter_paragraphs(document):
        for run in paragraph.runs:
            revised = normalize_run_text(run.text)
            if revised != run.text:
                run.text = revised
                changed_runs += 1

    document.save(args.output)
    print(f"Changed runs: {changed_runs}")


if __name__ == "__main__":
    main()
