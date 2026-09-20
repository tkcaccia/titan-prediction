from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path('/Users/stefano/Documents/Titan/titan-prediction/manuscript/JTM_full_audit_reviewer_report.docx')

NAVY = '17324D'
BLUE = '2B6F92'
TEAL = '2C7A7B'
PALE = 'EAF2F6'
PALE_TEAL = 'E9F4F2'
PALE_AMBER = 'FFF4DC'
AMBER = 'C27A12'
RED = 'A23B3B'
MID = '52606D'
LIGHT = 'D7E1E8'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    tr_pr.append(cant_split)


def set_table_borders(table, color=LIGHT, size='5'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color')
    c.set(qn('w:val'), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, size=None, bold=None, color=None, italic=None, font='Aptos'):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_labelled_paragraph(doc, label, text, style=None, color=NAVY):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    r = p.add_run(label)
    set_run_font(r, bold=True, color=color)
    p.add_run(text)
    return p


def add_callout(doc, title, text, fill=PALE_AMBER, accent=AMBER):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.75)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=175, bottom=150, end=175)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    set_run_font(r, size=8.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_metric_card(table, col, number, label, note=None):
    cell = table.cell(0, col)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE)
    set_cell_margins(cell, top=130, start=120, bottom=120, end=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(number)
    set_run_font(r, size=18, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(label)
    set_run_font(r2, size=8.5, bold=True, color=NAVY)
    if note:
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(note)
        set_run_font(r3, size=7.5, color=MID)


def add_major(doc, number, title, assessment, required, optional=None):
    h = doc.add_heading(f'{number}. {title}', level=2)
    h.paragraph_format.page_break_before = False
    add_labelled_paragraph(doc, 'Assessment. ', assessment)
    add_labelled_paragraph(doc, 'Required resolution. ', required, color=RED)
    if optional:
        add_labelled_paragraph(doc, 'Strengthening analysis. ', optional, color=TEAL)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.28)
sec.footer_distance = Inches(0.28)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
normal.font.size = Pt(9.4)
normal.font.color.rgb = RGBColor.from_string('263645')
normal.paragraph_format.space_after = Pt(5.2)
normal.paragraph_format.line_spacing = 1.06

for name, size, color, before, after in [
    ('Title', 25, NAVY, 0, 10),
    ('Heading 1', 17, NAVY, 12, 6),
    ('Heading 2', 12.5, BLUE, 9, 3),
    ('Heading 3', 10.5, TEAL, 7, 2),
]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for base, left, hanging in [('List Bullet', 0.28, 0.15), ('List Number', 0.30, 0.18)]:
    st = styles[base]
    st.font.name = 'Aptos'
    st.font.size = Pt(9.3)
    st.paragraph_format.left_indent = Inches(left)
    st.paragraph_format.first_line_indent = Inches(-hanging)
    st.paragraph_format.space_after = Pt(3.5)

if 'Small note' not in styles:
    small = styles.add_style('Small note', WD_STYLE_TYPE.PARAGRAPH)
else:
    small = styles['Small note']
small.font.name = 'Aptos'
small.font.size = Pt(8)
small.font.color.rgb = RGBColor.from_string(MID)
small.paragraph_format.space_after = Pt(3)

# Header and footer
header = sec.header
hp = header.paragraphs[0]
hp.text = 'CONFIDENTIAL PRE-SUBMISSION REVIEW  |  JOURNAL OF TRANSLATIONAL MEDICINE'
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in hp.runs:
    set_run_font(run, size=7.5, bold=True, color=BLUE)
footer = sec.footer
fp = footer.paragraphs[0]
add_page_number(fp)

# Memo masthead / cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run('JOURNAL-STYLE REVIEWER REPORT')
set_run_font(r, size=9, bold=True, color=TEAL)

title = doc.add_paragraph(style='Title')
title.add_run('Full manuscript and release audit')

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
r = sub.add_run('Major-revision assessment for Journal of Translational Medicine')
set_run_font(r, size=13, color=BLUE, italic=True)

meta = doc.add_table(rows=5, cols=2)
meta.autofit = False
meta.columns[0].width = Inches(1.42)
meta.columns[1].width = Inches(5.30)
set_table_borders(meta, color='C8D8E2', size='4')
labels = [
    ('Manuscript', 'Patient-level comparison of three released pathology embedding pipelines with a common PLS-based probe across 32 TCGA cancers'),
    ('Article/section', 'Research Article: Molecular Pathology'),
    ('Recommendation', 'Major revision; not suitable for acceptance in its current form'),
    ('Audit date', '25 August 2026'),
    ('Review basis', 'Main manuscript, 191-page Supplementary Material, response document, analysis/package code, machine-readable outputs and live release state'),
]
for i, (label, value) in enumerate(labels):
    prevent_row_split(meta.rows[i])
    for cell in meta.rows[i].cells:
        set_cell_margins(cell, top=105, start=120, bottom=105, end=120)
    set_cell_shading(meta.cell(i, 0), PALE)
    p1 = meta.cell(i, 0).paragraphs[0]
    r1 = p1.add_run(label)
    set_run_font(r1, size=8.5, bold=True, color=NAVY)
    p2 = meta.cell(i, 1).paragraphs[0]
    r2 = p2.add_run(value)
    set_run_font(r2, size=8.8, bold=(label == 'Recommendation'), color=(RED if label == 'Recommendation' else None))

doc.add_paragraph()
cards = doc.add_table(rows=1, cols=4)
cards.autofit = False
for c in range(4):
    cards.columns[c].width = Inches(1.67)
add_metric_card(cards, 0, '8,241', 'matched patients')
add_metric_card(cards, 1, '1,933', 'matched cancer-endpoint pairs')
add_metric_card(cards, 2, '323', 'TITAN permutation/FDR candidates')
add_metric_card(cards, 3, '83/323', 'lost original effect threshold', 'under TSS-code grouping')

doc.add_paragraph()
add_callout(
    doc,
    'Editorial bottom line',
    'The revision has a strong publishable core as a transparent, patient-level retrospective benchmark. Acceptance should depend on resolving the binary estimand, limiting claims to the tested probe and internal TCGA setting, elevating cohort-structure sensitivity, and completing a reproducible, licensable release. Independent validation would materially strengthen suitability for a translational journal.'
)

p = doc.add_paragraph(style='Small note')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('This report is written as an independent pre-submission review. It is not a legal opinion and does not claim to reproduce every model fit.')

doc.add_page_break()

doc.add_heading('Overall assessment', level=1)
p = doc.add_paragraph()
p.add_run('This is a substantially improved and unusually transparent computational pathology study. ').bold = True
p.add_run('The patient is consistently used as the unit of analysis and cross-validation; slide aggregation occurs before outcome matching; exact common-slide sensitivity addresses unequal embedding coverage; outcome provenance is audited; negative and ineligible results are preserved; and the manuscript now distinguishes the descriptive matched comparison from the deeper TITAN-only permutation/FDR screen. These choices make the work methodologically stronger than many TCGA-wide prediction surveys.')

p = doc.add_paragraph()
p.add_run('The central scientific contribution is nevertheless narrower than a translational prediction resource. ').bold = True
p.add_run('It is a retrospective comparison of three released embedding pipelines under a specific PLS-based linear probe, using internally resampled TCGA data and a catalogue dominated by correlated, computationally derived immune and genomic-context labels. The revised title reflects this limitation well. The manuscript should retain that hierarchy consistently and resist promoting threshold crossings, fitted-model objects or grouped internal validation as evidence of clinical transportability.')

p = doc.add_paragraph()
p.add_run('The most important unresolved methodological issue is the binary estimand. ').bold = True
p.add_run('Latent-component selection and crossing status are driven by balanced accuracy from empirical-prior LDA class calls, yet alternative equal-prior and training-threshold rules nearly double some crossing counts. The choice therefore changes representation consensus and downstream resource membership, rather than merely changing a secondary display metric. This must be resolved before the atlas can be interpreted confidently.')

doc.add_heading('Principal strengths', level=1)
strengths = [
    ('Patient-first validation.', 'Multiple slides from one participant are pooled before outcome matching, and participants rather than slides enter validation folds.'),
    ('Fair matched representation inputs.', 'The common-patient analysis is supplemented by an exact common-slide audit: 8,207 of 8,241 patients had identical slide sets, with 10,165 shared slides, and the common-slide sensitivity indicates that aggregate differences are not explained by missing slides.'),
    ('Explicit evidence layers.', 'The manuscript separates 1,933 matched descriptive tasks from the 2,073-task TITAN screen with permutation and multiplicity control, and it no longer presents their evidential standards as equivalent.'),
    ('Provenance-aware endpoints.', 'Direct genomic alterations, sequencing-derived burdens, inferred cell fractions, transcriptomic signatures, composite scores and same-H&E quantities are distinguished.'),
    ('Cohort-structure sensitivity.', 'Tissue-source-site-code grouping is conducted in both inner and outer validation, supplemented by matched-random controls and code-only prediction analyses.'),
    ('Reproducibility orientation.', 'The analysis records schemas, seeds, feature order, component choices, registry fields, hashes, fold adequacy, class counts and model availability status.'),
    ('Honest limitations.', 'The manuscript acknowledges internal-only validation, downstream-probe dependence, uncertain model redistribution, lack of calibration and absence of pathology-adjudicated quality control.'),
]
for label, body in strengths:
    add_labelled_paragraph(doc, label + ' ', body, color=TEAL)

doc.add_heading('Major concerns and required revisions', level=1)

add_major(
    doc, 1, 'Align the binary tuning objective, primary metric and crossing rule',
    'The primary binary pipeline selects 1 to 10 PLS components using inner out-of-fold balanced accuracy from empirical-training-prior LDA calls and defines crossings by balanced accuracy ≥0.60. AUROC and precision-recall AUC are reported, but they are generated by a component count chosen under the empirical-prior balanced-accuracy objective. The operating rule is therefore integral to the estimated representation performance rather than a cosmetic threshold choice. The atlas-wide sensitivity is decisive: TITAN binary crossings change from 93 with empirical priors to 193 with equal priors and 184 with an inner-optimised threshold, with corresponding shifts for the other representations.',
    'Choose and declare one coherent binary estimand. The preferred benchmark for representation comparison is AUROC-centred: tune component count by pooled inner out-of-fold AUROC; use outer out-of-fold AUROC as the primary paired statistic; and report precision-recall AUC plus balanced accuracy, sensitivity, specificity, PPV and NPV at a strictly training-derived operating point. Alternatively, retain empirical-prior PLS-LDA as the primary probe, but then state that balanced accuracy under that exact rule is the primary estimand. AUROC should not be described as primary in that case, and crossing or consensus labels should not be treated as general representation properties. Any change must regenerate binary crossings, consensus summaries, grouped-retention labels and the model registry.',
    'Report the full atlas-wide comparison of empirical priors, equal priors and training-only threshold optimisation as continuous paired metric changes in addition to crossing-count changes. Include prevalence and the no-skill precision-recall reference for highlighted binary outcomes.'
)

add_major(
    doc, 2, 'Keep the central result probe-specific and resolve the component ceiling',
    'The paper correctly shows that the apparent representation leader depends on the downstream method. Ridge retains the PLS-defined leader in only 17 of 35 binary and 8 of 12 continuous comparisons. At least one outer fit reaches the ten-component ceiling in 42.5% of Giga-SSL binary tasks, 35.9% of Prov-GigaPath binary tasks and 38.5% of TITAN binary tasks; widening the range changes several winners in the 47-target subset. These are not peripheral observations when the central claim is a representation comparison.',
    'Either extend the wider-component and/or symmetric ridge sensitivity across the 426 matched binary tasks, or at minimum across every union-positive and near-threshold binary pair, using identical folds and tuning rules. If this is computationally infeasible, narrow every conclusion to “performance under the fixed 1 to 10 component PLS/PLS-LDA probe.” Do not describe the findings as intrinsic foundation-model superiority.',
    'For each representation, report selected-component distributions, ceiling-hit frequency, near-constant feature handling, numerical failures and fallback behaviour. A paired comparison of TITAN and Prov-GigaPath is particularly interpretable because their reported pretraining corpora exclude TCGA.'
)

add_major(
    doc, 3, 'Treat the matched atlas as descriptive unless representation-specific inference is added',
    'The 1,933-task matched benchmark reports performance-threshold crossings, not representation-specific permutation/FDR discoveries. The TITAN-only layer applies a different evidential standard and task universe. This distinction is now stated, but raw crossing totals and consensus categories still risk being read as discovery counts, especially because correlated Thorsson-derived outcomes dominate the continuous catalogue.',
    'Keep “effect-threshold crossing” throughout the matched comparison; never shorten it to “positive,” “discovery” or “screen-positive.” Lead with paired Q²/AUROC distributions, macro-averaged cancer and endpoint-family summaries, unique endpoint definitions and crossing proportions. State prominently that task-level counts are catalogue-dependent and are not independent biological discoveries. If inferential representation claims are desired, add representation-specific permutation/multiplicity procedures on a prospectively identified restricted target set.',
    'For union-positive and near-threshold tasks, report crossing frequency, paired effect variability and consensus-class stability over common alternative partitions. Emphasise ranks and effect differences rather than hard categories near Q²=0.20 or balanced accuracy=0.60.'
)

add_major(
    doc, 4, 'Independent validation remains the central translational limitation',
    'All estimates remain internal to TCGA. Nested cross-validation, repeat partitions, exact common-slide analysis and tissue-source-site-code grouping reduce specific biases but cannot establish transportability to other institutions, scanners, staining distributions, populations or molecular-assay workflows. For Journal of Translational Medicine, this materially limits any implication of clinical prediction or deployment.',
    'Preferably, externally evaluate a locked and explicitly prespecified subset spanning one common mutation, one MSI/genomic-instability target and one continuous sequence- or transcriptome-derived phenotype. Use the documented representation pipeline without refitting, threshold adjustment or selective omission, and report failures. If compatible external data are unavailable, retain the current benchmark/resource framing, describe all estimates as “internally derived TCGA estimates,” and use “research software demonstration” rather than “deployment,” “patient profile” or assay-replacement language.',
    'A small, honest locked external test is more valuable than adding further internal sensitivity analyses. External failure would still be informative if all targets are reported.'
)

add_major(
    doc, 5, 'Make upstream training exposure and pipeline differences central to interpretation',
    'The three resources differ in pretraining data, tile and slide encoders, dimensionality, spatial resolution, preprocessing and released layer. Giga-SSL development used TCGA, so direct overlap between representation-learning images and evaluated slides cannot be excluded. Downstream molecular labels are held out, but the image-exposure condition differs from TITAN and Prov-GigaPath. The study therefore compares released embedding pipelines, not isolated foundation-model quality.',
    'State the Giga-SSL exposure qualification in the Abstract and principal comparison table, not only in Methods/limitations. Add explicit inclusion criteria for the three resources and a short supplementary inventory of relevant models excluded because reproducible slide-level TCGA embeddings, identifiers or release metadata were unavailable. Use “released representation pipeline” consistently and avoid model-superiority language.',
    'Present a separate paired TITAN and Prov-GigaPath summary so readers can interpret the two pipelines whose reported pretraining corpora excluded TCGA.'
)

add_major(
    doc, 6, 'Integrate tissue-source-site-code sensitivity into the primary evidence display',
    'Eighty-three of 323 TITAN candidates fall below their original effect threshold when complete TCGA tissue-source-site codes are held apart. READ-APC and COAD-APC fall close to chance, and tissue-source-site code alone predicts APC status. This is a central cohort-structure result. However, grouped folds can be sparse or single-class, and the barcode-derived code is not an institution, scanner, laboratory or staining-batch identifier. Grouped attenuation therefore supports sensitivity to cohort structure, not proof of technical confounding or transportability.',
    'Show the grouped estimate beside every highlighted random-fold estimate and retain the matched-random comparator. Explicitly report whether metrics pool all outer-fold predictions, how single-class folds are handled, the number of affected binary tasks, minimum inner/outer training-class counts, number of contributing codes and whether four or five folds were used. Add a machine-readable fold-adequacy flag and prevent sparse grouped estimates from driving an unqualified robustness label.',
    'Apply tissue-source-site-code grouping to Giga-SSL and Prov-GigaPath for the common-cohort union-positive set if feasible. Otherwise state prominently that grouped sensitivity is TITAN-specific and that multi-representation consensus is not evidence of robustness to cohort structure.'
)

add_major(
    doc, 7, 'Report uncertainty as conditional and reduce dependence on threshold labels',
    'Repeated nested validation is performed after selection. The reported patient-resampling intervals reuse fixed held-out predictions and therefore do not repeat screening, fold generation, tuning or fitting. The manuscript now calls these selection-conditioned intervals, which is appropriate, but readers still need a compact statement of what variability is and is not represented. Hard crossing categories remain unstable around pragmatic thresholds.',
    'In each highlighted table, label the primary-screen estimate and five-repeat mean separately and use the full term “selection-conditioned patient-resampling interval for repeated out-of-fold predictions” in the note. Report between-partition distributions, paired representation differences and crossing proportions. Do not present these intervals as confidence intervals for external generalisation or as correction for winner’s-curse selection.',
    'The permutation audit is otherwise strong: checkpoint failures and early-stopped tests receive distinct p=1 statuses; complete modelling is repeated; and selected targets are refined to 9,999 permutations with Monte Carlo intervals. Keep those details in the Supplement and summarize them once in the main Methods.'
)

add_major(
    doc, 8, 'Elevate pathology quality control, molecular-slide linkage and pooling limitations',
    'Patient-level pooling correctly prevents a participant with many slides from receiving greater validation weight, but equal mean pooling cannot ensure that every slide contains representative tumour. No structured tumour cellularity, tissue area, artefact, biopsy/resection or slide-quality fields were available. Molecular outcomes may link only by participant or sample barcode, not the same block, analyte or tumour region. The generated narrative audit identified six patients with “no residual tumour” language; this is not an adjudicated pathology label, but it illustrates the risk.',
    'Add a concise main-text schematic or table describing, by endpoint source, whether linkage is at participant, sample or aliquot level and what spatial/aliquot mismatch remains. Report the six affected patients and whether highlighted results materially change when they are excluded, clearly labelling this as non-adjudicated sensitivity. State that exact sample-barcode agreement does not establish block- or region-level identity.',
    'Using existing embeddings, add mean-versus-median pooling, within-patient cosine dispersion and a sensitivity excluding the 30-slide SARC participant. Leave-one-slide-out analysis for multi-slide patients would further quantify pooling stability without new annotation.'
)

add_major(
    doc, 9, 'Separate direct genomic outcomes from derived and same-H&E reference phenotypes',
    'The provenance audit is excellent, but approximately 93% to 97% of continuous crossings arise from the Thorsson-derived catalogue. These labels include computationally inferred immune-cell fractions, transcriptomic signatures, methylation-derived quantities, repertoire measures, in-silico neoantigens, composite genomic-context scores and an H&E-derived TIL fraction. Agreement with these labels is not equivalent to recovering a directly measured immune-cell count or a certified biomarker.',
    'Lead the Abstract and principal results with provenance-stratified rates: direct genomic alterations, sequencing-derived burdens, inferred immune fractions, transcriptomic signatures, composite scores and same-H&E quantities. Use same-H&E-excluded values as the default cross-modal continuous summary. Describe CIBERSORT and RNA-signature results as agreement with computational reference phenotypes, and do not use “molecular prediction” as an unqualified umbrella term.',
    'For a small number of strong, consensus, grouped-stable examples, add blinded pathologist review or spatial relevance/localisation. Nearest-neighbour examples alone show similarity, not the morphology supporting a prediction.'
)

add_major(
    doc, 10, 'Clarify the fitted-object universe, access route and upstream licensing',
    'The local PathoFMPred registry contains 917 objects, but the Giga-SSL and Prov-GigaPath objects represent a shared target universe derived from TITAN-qualified matched-eligible targets, not a complete set of each alternative representation’s own crossings. Fitted objects remain access controlled. A model-free public shell can validate schemas and demonstrate interfaces but cannot deliver the inference functionality emphasised in the resource claim. TITAN-derived redistribution also remains subject to upstream permission; no legal conclusion should be implied.',
    'State exactly which target universe each representation’s objects cover, whether an object may exist when that representation does not cross its threshold, and what compare_pathofm_models() returns in that circumstance. If the alternative-representation-specific atlas is not fully operationalised, describe PathoFMPred as an “analysis and inference framework with access-controlled reference models,” not a complete reusable multi-representation model resource. Distinguish the GPL-3 source-code licence from fitted-object access terms and each upstream representation’s derivative-use terms. Provide a durable controlled-access process or reduce the inference claim accordingly.',
    'The package itself has a strong technical baseline: the available PathoFMPred 0.2.0 check log completes R CMD check --as-cran with only the standard new-submission NOTE. Preserve the vignette, reference manual, synthetic fixtures and schema validation in the public shell.'
)

add_major(
    doc, 11, 'Archive one synchronized, immutable submission snapshot',
    'The document audit scripts pass and many previous cross-reference inconsistencies have been corrected. However, the live analysis repository remains a working tree with 236 modified/untracked paths relative to commit 01dab393f32858253dd33f54f5f762ada40e0e9d, and the PathoFMPred working tree also contains extensive changes. The manuscript still has placeholders for final release tag, immutable commit and DOI. A reader cannot reconstruct the submitted result from a moving private tree.',
    'Regenerate the manuscript, supplement, figures, registries and software documentation from one clean results snapshot; commit that snapshot; create a versioned release; archive the analysis code and model-free package shell with a DOI; and insert the exact tag, commit and DOI into the manuscript. The deposited machine-readable files should be described as complete-resolution companions, not as an authority competing with the paper. Include a machine-readable manifest of file hashes and analysis versions.',
    'Before submission, rerun document cross-reference checks, package tests and R CMD check from the archived source rather than the current working directory, and preserve the logs in the release.'
)

add_major(
    doc, 12, 'Complete the editorial package and substantially reduce the Supplement',
    'The scientific narrative is much clearer, but the 191-page Supplement remains difficult to navigate. The main document contains only three figures and two numbered tables, which is appropriately lean, yet the Supplement contains 76 tables and 13 figures. The reference audit also indicates that only references 1 to 16, 30, 31, 39 and 40 are cited in the manuscript or supplement text; 22 of 42 bibliography entries appear uncited. Competing interests, funding and author contributions remain submission placeholders.',
    'Complete all declarations before submission. Audit every reference and either cite it at the relevant Methods/endpoint-source/literature statement or remove it. Verify the spelling of the author name in reference 35 against the canonical journal record. Add a concise Supplement table of contents, retain interpretive summary tables and figures in the PDF, and move exhaustive per-model records to machine-readable files. Ensure all additional files are present and named exactly as cited.',
    'In Figure 1, “11,449 eligible slides” is the TITAN release count within a three-pipeline workflow; label it explicitly as TITAN or show all three released counts plus the exact common-slide count. Keep representation colours and order identical throughout and preserve readable font sizes at final journal width.'
)

doc.add_heading('Prior concerns now substantially addressed', level=1)
resolved = [
    ('Study hierarchy and title', 'The current title accurately describes a comparison of released embedding pipelines using a common PLS-based probe, and the analysis map separates matched, TITAN inferential and software layers.'),
    ('Histological input matching', 'The exact common-slide audit resolves the main concern that patient matching might conceal different slide subsets.'),
    ('Crossing terminology', 'The manuscript now distinguishes descriptive effect-threshold crossings from permutation/FDR-qualified TITAN candidates.'),
    ('Multiplicity mechanics', 'Below-checkpoint and early-stopped models receive distinct p=1 statuses, remain in Benjamini-Hochberg denominators, and leading claims receive 9,999-permutation precision refinement with Monte Carlo intervals.'),
    ('Evidence-maturity audits', 'Small binary classes and smaller continuous models are separately labelled, with fold counts, component distributions and stability analyses reported.'),
    ('Chronology', 'The study is described as retrospective with documented analysis settings; prospectively locked language has largely been removed.'),
    ('Site terminology', '“TCGA tissue-source-site code” is used instead of implying institution- or scanner-level validation.'),
    ('Software naming', 'The manuscript and package now consistently use PathoFMPred, with representation-specific input schemas.'),
]
for label, text in resolved:
    add_labelled_paragraph(doc, label + '. ', text, color=TEAL)

doc.add_heading('Prioritized revision matrix', level=1)
table = doc.add_table(rows=1, cols=4)
table.autofit = False
widths = [1.45, 2.90, 1.08, 1.28]
for i, w in enumerate(widths):
    table.columns[i].width = Inches(w)
headers = ['Issue', 'Minimum acceptable resolution', 'New fitting?', 'Editorial weight']
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=105, start=90, bottom=105, end=90)
    p = cell.paragraphs[0]
    r = p.add_run(h)
    set_run_font(r, size=8, bold=True, color=WHITE)
set_repeat_table_header(table.rows[0])
rows = [
    ('Binary estimand', 'Align tuning, primary metric and operating rule; regenerate categorical outputs if the rule changes.', 'Yes', 'Decisive'),
    ('Probe/component sensitivity', 'Full binary atlas or all union-positive/near-threshold tasks under wider grid or matched ridge.', 'Yes', 'Decisive'),
    ('External validation', 'Locked external subset, or explicitly internal benchmark framing with no deployment claims.', 'Preferred', 'High'),
    ('TSS-code sensitivity', 'Main tables with grouped/matched-random estimates plus fold-adequacy reporting.', 'Limited', 'High'),
    ('Pathology/linkage', 'Main-text linkage/QC summary and numerical pooling sensitivities.', 'Limited', 'High'),
    ('Endpoint provenance', 'Provenance-stratified headline summaries; same-H&E excluded by default.', 'No', 'High'),
    ('Software/licensing', 'Exact object universe, durable access terms and upstream redistribution resolution.', 'No', 'Blocker for resource claim'),
    ('Immutable release', 'Clean synchronized commit, versioned archive, DOI and checksums.', 'No', 'Blocker for reproducibility'),
    ('Editorial package', 'Declarations, reference audit and condensed navigable Supplement.', 'No', 'Required'),
]
for row_i, data in enumerate(rows, start=1):
    cells = table.add_row().cells
    prevent_row_split(table.rows[row_i])
    for col_i, val in enumerate(data):
        set_cell_margins(cells[col_i], top=90, start=90, bottom=90, end=90)
        if row_i % 2 == 0:
            set_cell_shading(cells[col_i], 'F5F8FA')
        p = cells[col_i].paragraphs[0]
        r = p.add_run(val)
        set_run_font(r, size=7.7, bold=(col_i == 0 or col_i == 3), color=(RED if col_i == 3 and 'Decisive' in val else None))
set_table_borders(table, color='C5D3DC', size='4')

doc.add_heading('Minor and editorial comments', level=1)
minor = [
    'Define Q², q-value, out-of-fold (OOF), selection-conditioned (SC) interval and tissue-source-site code at first use as well as in the abbreviation list.',
    'Where initial-screen and five-repeat estimates both appear, label them in the same row or caption; do not allow readers to interpret the two values as an inconsistency.',
    'Retain “sample-size maturity stratum” rather than “standard evidence” and use threshold categories only as database navigation tags.',
    'Keep the two COAD radar cases as a software illustration, but place them in the Supplement, state their post hoc selection, show original predicted values, and place “reference rank, not probability” immediately beside any binary rank.',
    'Do not narrate treatment response for the two COAD cases; treatment was neither an input nor a validation outcome.',
    'Use Giga-SSL, Prov-GigaPath, PathoFMPred, cancer-endpoint pair and tissue-source-site code consistently.',
    'For any standardized LDA score plot, a zero line must be labelled only as a visual reference unless it is mathematically the fold-specific decision threshold after transformation.',
    'Keep implementation details such as model hashes, score-rank construction and report-field semantics in software documentation or machine-readable metadata rather than the biological Results.',
    'Ensure the literature crosswalk is labelled as a targeted narrative audit unless database search strings, dates, inclusion criteria and mapping rules are fully reported.',
    'The current journal fit is plausible in Medical Bioinformatics, Molecular Pathology, Disease Biomarkers and Translational Imaging, but the cover letter should present the work as a reproducible translational research prioritisation resource. It should not present the study as clinical validation.'
]
for item in minor:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

doc.add_heading('Confidential comments to the editor', level=1)
p = doc.add_paragraph()
p.add_run('This manuscript is more rigorous and transparent than many large TCGA prediction screens. ').bold = True
p.add_run('Patient-level aggregation, exact common-slide sensitivity, preservation of negative results, endpoint provenance and tissue-source-site-code analyses are genuine strengths. The authors have also corrected many earlier presentation and reproducibility inconsistencies.')
p = doc.add_paragraph()
p.add_run('I nevertheless recommend major revision. ').bold = True
p.add_run('The primary binary conclusions depend strongly on the LDA prior/operating rule, a large fraction of binary fits reaches the component ceiling, and the representation ranking is visibly downstream-probe dependent. There is no independent validation, while some prominent TITAN associations attenuate to near chance under tissue-source-site-code grouping. The fitted-model resource is access controlled and not yet tied to an immutable, licensable public release.')
p = doc.add_paragraph()
p.add_run('I would support publication if the binary estimand is resolved, the claims remain explicitly probe- and TCGA-specific, cohort-structure sensitivity is integrated into the main evidence display, and the analysis/software snapshot is archived reproducibly. ').bold = True
p.add_run('External validation would materially strengthen suitability for Journal of Translational Medicine; without it, the article should be judged as a computational pathology benchmark and research resource rather than a translational diagnostic study.')

doc.add_heading('Summary judgement', level=1)
summary = doc.add_table(rows=1, cols=2)
summary.autofit = False
summary.columns[0].width = Inches(2.25)
summary.columns[1].width = Inches(4.48)
set_table_borders(summary, color='C5D3DC', size='4')
for c, h in enumerate(['Criterion', 'Assessment']):
    set_cell_shading(summary.cell(0, c), NAVY)
    r = summary.cell(0, c).paragraphs[0].add_run(h)
    set_run_font(r, size=8.5, bold=True, color=WHITE)
set_repeat_table_header(summary.rows[0])
judgements = [
    ('Scientific question', 'Important and timely'),
    ('Patient-level design', 'Strong'),
    ('Matched representation audit', 'Strong and substantially improved'),
    ('Binary estimand', 'Requires major statistical revision'),
    ('Probe independence', 'Insufficient; claims must remain probe-specific'),
    ('Biological interpretation', 'Moderate; dominated by derived reference phenotypes'),
    ('External/translational validation', 'Insufficient'),
    ('Transparency', 'Very strong'),
    ('Software/release readiness', 'Technically promising but access/licensing/archive incomplete'),
    ('Presentation', 'Main text improved; Supplement overly long'),
    ('Current recommendation', 'Major revision'),
]
for idx, (criterion, assessment) in enumerate(judgements, start=1):
    cells = summary.add_row().cells
    prevent_row_split(summary.rows[idx])
    for j, value in enumerate((criterion, assessment)):
        set_cell_margins(cells[j], top=85, start=105, bottom=85, end=105)
        if idx % 2 == 0:
            set_cell_shading(cells[j], 'F5F8FA')
        r = cells[j].paragraphs[0].add_run(value)
        set_run_font(r, size=8.2, bold=(j == 0 or criterion == 'Current recommendation'), color=(RED if criterion == 'Current recommendation' else None))

doc.add_heading('Audit scope and verification note', level=1)
p = doc.add_paragraph()
p.add_run('Inspected materials. ').bold = True
p.add_run('Main manuscript, Supplementary Material, response-to-reviewer document, analysis scripts, result registries, package source/check logs, Git working-tree and repository visibility state. Automated document and release audits completed successfully for the reported task/model counts. The available PathoFMPred 0.2.0 R CMD check --as-cran log completed with one new-submission NOTE and no warning or error.')
p = doc.add_paragraph()
p.add_run('Limits of this review. ').bold = True
p.add_run('I did not refit all 1,933 matched tasks or independently reproduce every 9,999-permutation result. I did not perform a legal determination concerning derivative-model redistribution. Numerical statements above are therefore an evidence-backed manuscript/release audit, not a full independent replication.')

# Document metadata
props = doc.core_properties
props.title = 'Full JTM manuscript and release audit'
props.subject = 'Journal-style reviewer report for a pathology foundation-model benchmark'
props.author = 'Independent pre-submission review'
props.keywords = 'Journal of Translational Medicine; reviewer report; computational pathology; foundation models; TCGA; PLS'

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
