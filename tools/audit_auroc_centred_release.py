#!/usr/bin/env python3
"""Fail closed on the AUROC-centred PathoFMPred manuscript release."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
MAIN = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
SUPP = ROOT / "manuscript" / "supplementary_material_JTM.docx"
RESPONSE = ROOT / "manuscript" / "response_to_reviewer_JTM.docx"
COVER = ROOT / "manuscript" / "cover_letter_JTM.docx"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def full_text(path: Path) -> tuple[Document, str]:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return doc, "\n".join(chunks)


main, main_text = full_text(MAIN)
supp, supp_text = full_text(SUPP)
response, response_text = full_text(RESPONSE)
cover, cover_text = full_text(COVER)
combined = "\n".join((main_text, supp_text, response_text, cover_text))

require(main_text.startswith(
    "Patient-level comparison of three released pathology representation pipelines "
    "with a common PLS-based probe across 32 TCGA cancers"
), "Current title is missing")
for phrase in (
    "Pathology foundation models are neural networks pretrained",
    "In this study, a cancer-endpoint pair denotes one TCGA cancer type",
    "We also created PathoFMPred as a reproducible analysis interface",
    "The matched representation comparison, patient-level analysis and transparent reporting",
):
    require(phrase in main_text, f"Introduction is missing: {phrase}")

background = main_text.split("Background\n", 1)[1].split("\nMethods\n", 1)[0]
citations = re.findall(r"\[(\d+)", background)
require(citations and citations[0] == "1", "The first Background citation is not reference 1")

for forbidden in ("—", "–", "TITANPred", "revision-added", "PLS1", "PLS2"):
    require(forbidden not in combined, f"Forbidden or obsolete term remains: {forbidden}")
require("?" not in main_text, "A question mark remains in the main manuscript")

for phrase in (
    "pooled inner out-of-fold (OOF) AUROC selected the component count",
    "outer OOF AUROC was the primary statistic",
    "A training-only threshold supplied balanced accuracy, sensitivity, specificity, PPV and NPV",
    "PR-AUC was reported with prevalence as its no-skill reference",
    "Q²≥0.20 or AUROC≥0.60",
):
    require(phrase in main_text, f"Binary estimand is incompletely stated: {phrase}")

for phrase in (
    "237/170/180",
    "86/186/169 for TITAN",
    "62/115/117 for Giga-SSL",
    "79/113/124 for Prov-GigaPath",
    "0.940/0.878/0.873",
    "0.805/0.653/0.682",
    "no-skill PR-AUC reference of 0.150",
):
    require(phrase in main_text, f"Required primary or operating-rule result is missing: {phrase}")

for phrase in (
    "median balanced accuracy by 0.030/0.021/0.017",
    "sensitivity by 0.378/0.381/0.365",
    "specificity by -0.303/-0.329/-0.328",
    "PPV by -0.020/-0.012/-0.020",
    "NPV by 0.019/0.012/0.013",
):
    require(phrase in main_text, f"Paired continuous operating-rule changes are missing: {phrase}")

figure_captions = [
    p.text for p in main.paragraphs if re.match(r"^Figure \d+\. ", p.text)
]
require([int(re.match(r"Figure (\d+)", x).group(1)) for x in figure_captions]
        == [1, 2, 3], "Main figures are not exactly 1 to 3 in order")
require("Patient and slide overlap across the three released TCGA representation datasets" in figure_captions[0],
        "Figure 1 is not the patient-overlap display")
for phrase in (
    "11,449 slides from 9,404 patients",
    "11,427 slides from 9,378 patients",
    "10,328 slides from 8,393 patients",
    "10,165 slides",
    "8,207 patients had identical slide sets",
):
    require(phrase in figure_captions[0], f"Figure 1 caption is missing: {phrase}")
require("PathoFMPred research-software profiles for two COAD patients" not in main_text,
        "The post hoc COAD profiles remain in the main manuscript")
table_captions = [p.text for p in main.paragraphs if re.match(r"^Table \d+\. ", p.text)]
require(len(table_captions) == 3, "The main manuscript should contain three tables")
require(table_captions[0].startswith(
    "Table 1. Released representation pipelines and upstream training exposure"
), "Table 1 does not place upstream exposure in the principal comparison")
require(table_captions[1].startswith(
    "Table 2. Tumour-feature classes with descriptive effect-threshold crossings"
), "Table 2 is not the tumour-feature summary")
require(table_captions[2].startswith(
    "Table 3. Partition and cohort-structure sensitivity for highlighted binary cancer-endpoint pairs"
), "Table 3 does not place partition and grouped estimates beside highlighted primary estimates")
for phrase in (
    "Five alternative partitions: median [IQR]; crossing proportion",
    "These distributions and descriptive warnings are internal sensitivity summaries",
    "Median alternative-minus-TITAN differences",
    "These cancer-cluster intervals describe paired task effects and are not confidence intervals for external generalisation",
):
    require(phrase in main_text, f"Conditional partition presentation is missing: {phrase}")

for phrase in (
    "direct overlap between representation-learning images and evaluated slides cannot be excluded",
    "the reported TITAN and Prov-GigaPath pretraining corpora excluded TCGA",
    "We applied five pragmatic inclusion criteria",
    "Supplementary Table S15e lists representative alternatives excluded",
    "both pipelines crossed Q² 0.20 in 121/1,507 pairs",
    "both crossed AUROC 0.60 in 157/426 pairs",
    "still compares complete pipelines",
):
    require(phrase in main_text, f"Upstream-exposure interpretation is missing: {phrase}")
require("released embedding pipeline" not in main_text,
        "Obsolete released-embedding-pipeline terminology remains in the manuscript")

for phrase in (
    "paired Q² and AUROC values without thresholding",
    "Catalogue-normalised summaries",
    "not independent biological discoveries",
    "The matched atlas has no representation-specific permutation or multiplicity testing",
    "165/322, 80/322 and 102/322 continuous pairs",
    "186/366, 119/366 and 133/366",
    "146/322 and 136/366",
):
    require(phrase in main_text, f"Descriptive matched-atlas presentation is missing: {phrase}")
require("union-positive" not in main_text,
        "Obsolete union-positive terminology remains in the main manuscript")
require("Paired target-level performance across all 1,933 matched cancer-endpoint tasks"
        in figure_captions[1], "Figure 2 does not lead with paired effects")

for phrase in (
    "Aamilah Ismail3,*, Martin Ocharo1,2,*",
    "Silvano Piazza6,8, Dinesh Gupta7, Stefano Cacciatore1,2",
):
    require(phrase in main_text, f"Required author identifier is missing: {phrase}")
for phrase in ("TCGA-AA-A01F", "TCGA-A6-A56B"):
    require(phrase in supp_text, f"Required supplementary case identifier is missing: {phrase}")

reference_block = main_text.split("\nReferences\n", 1)[1]
reference_numbers = [int(x) for x in re.findall(r"(?m)^(\d+)\. ", reference_block)]
require(reference_numbers == list(range(1, 43)), "References are not numbered 1 to 42 in order")
require("Wells K, Lamrca A, Papaxoinis G, Wallace A, Quinn AM, Summers Y, Nonaka D" in
        reference_block, "Reference 35 does not match the canonical author record")

scientific_text = main_text.split("\nReferences\n", 1)[0] + "\n" + supp_text
cited_numbers: set[int] = set()
for match in re.finditer(r"\[([0-9, -]+)\]", scientific_text):
    for token in match.group(1).split(","):
        token = token.strip()
        if "-" in token:
            start, end = (int(value) for value in token.split("-", 1))
            cited_numbers.update(range(start, end + 1))
        elif token:
            cited_numbers.add(int(token))
require(cited_numbers == set(range(1, 43)),
        f"Uncited or invalid retained references: {sorted(set(range(1, 43)) - cited_numbers)}")

supplement_figure_captions = [
    p.text for p in supp.paragraphs if re.match(r"^Figure S\d+\. ", p.text)
]
require([int(re.match(r"Figure S(\d+)", caption).group(1))
         for caption in supplement_figure_captions] == list(range(1, 8)),
        "Supplementary figures are not exactly S1 to S7 in order")
require(len(supp.inline_shapes) == 7, "The Supplement must contain seven figures")
require("Figure S7. Post hoc PathoFMPred research-software illustration" in supp_text,
        "The COAD radar is not Supplementary Figure S7")
require("Reference rank, not probability" in supp_text,
        "The supplementary software illustration lacks the binary-rank warning")
require(len(supp.tables) < 76, "The Supplement still contains too many embedded tables")
for phrase in (
    "supplementary_material_JTM.docx",
    "Additional_file_2_COAD_example_A_PathoFMPred_report.pdf",
    "Additional_file_3_COAD_example_B_PathoFMPred_report.pdf",
):
    require(phrase in main_text and phrase in supp_text,
            f"Additional-file inventory is not synchronized: {phrase}")
for filename in (
    "Additional_file_2_COAD_example_A_PathoFMPred_report.pdf",
    "Additional_file_3_COAD_example_B_PathoFMPred_report.pdf",
):
    require((ROOT / "manuscript" / filename).is_file(),
            f"Cited additional file is missing: {filename}")
require("12. Complete the editorial package and substantially reduce the Supplement"
        in response_text, "The latest editorial-package response is missing")
require("Competing interests\nPending corresponding-author confirmation" in main_text and
        "Funding\nPending corresponding-author confirmation" in main_text and
        "Authors' contributions\nPending approval by all authors" in main_text,
        "Author-dependent declarations are not explicitly flagged")

require("Table S10j. AUROC-centred matched binary benchmark" in supp_text,
        "Supplementary AUROC-centred benchmark table is missing")
require("Align the binary tuning objective, primary metric and crossing rule" in response_text,
        "The latest reviewer comment is missing from the response")
require("We adopted the requested AUROC-centred benchmark" in response_text,
        "The response does not state the implemented resolution")
for phrase in (
    "323/366 (88.3%)",
    "changed 43",
    "89/426 TITAN",
    "98/426 Giga-SSL",
    "126/426 Prov-GigaPath",
    "without a numerical failure or fallback",
):
    require(phrase in main_text, f"Expanded-component result is missing: {phrase}")
require("2. Keep the central result probe-specific and resolve the component ceiling"
        in response_text, "The latest component-ceiling response is missing")
require("All 5,490 expanded-grid outer fits completed with no numerical failure or fallback"
        in response_text, "The response omits the expanded-grid failure audit")
require("3. Treat the matched atlas as descriptive unless representation-specific inference is added"
        in response_text, "The latest descriptive-atlas response is missing")
require("did not add a post hoc representation-specific permutation/FDR analysis"
        in response_text, "The response does not state the inferential decision")
require("AUROC thresholds 0.55-0.65" in supp_text,
        "Supplementary threshold sensitivity is not labelled with AUROC")
require("Table S15e. Representative pathology foundation models not included"
        in supp_text, "The supplementary exclusion inventory is missing")
require("Table S15f. Dedicated paired TITAN-Prov-GigaPath summary"
        in supp_text, "The dedicated TITAN-Prov-GigaPath summary is missing")
require("5. Make upstream training exposure and pipeline differences central to interpretation"
        in response_text, "The upstream-exposure reviewer response is missing")
require("New main Table 1 places this qualification" in response_text,
        "The response does not identify the principal exposure table")

for phrase in (
    "Of 1,488 representation-task fits, 1,485 were estimable",
    "ACC genome doubling was not estimable for any representation",
    "28 had a single-class outer test fold",
    "120 had a single-class inner validation fold",
    "162/273 binary and 51/223 continuous tasks",
    "83/323 candidates (25.7%) fell below their original threshold",
    "READ-APC declined from balanced accuracy 0.862 to 0.495",
    "COAD-APC from 0.793 to 0.543",
    "not proof of technical confounding or transportability",
):
    require(phrase in main_text, f"Primary cohort-structure result is missing: {phrase}")
require("6. Integrate TCGA tissue-source-site code sensitivity into the primary evidence display"
        in response_text, "The current tissue-source-site reviewer response is missing")
require("Main Figure 3 now places TCGA tissue-source-site-code-grouped Q² or AUROC"
        not in response_text,
        "Hyphenated tissue-source-site-code terminology remains in the response")

require("7. Report uncertainty as conditional and reduce dependence on threshold labels"
        in response_text, "The conditional-uncertainty reviewer response is missing")
for phrase in (
    "95% selection-conditioned patient-resampling interval for repeated out-of-fold predictions",
    "does not correct winner's-curse selection",
    "not a confidence interval for external generalisation",
    "repeat median",
    "crossing proportion",
):
    require(phrase in supp_text, f"Highlighted uncertainty presentation is missing: {phrase}")

for phrase in (
    "Box 1. Pathology-quality, pooling and molecular-slide linkage constraints",
    "TCGA-B6-A0IA and TCGA-B6-A0WV in BRCA",
    "TCGA-55-8203 and TCGA-55-8507 in LUAD",
    "TCGA-22-4596 in LUSC and TCGA-DX-AB2L in SARC",
    "64/69 models",
    "SARC Macrophage Regulation Q² from 0.520 to 0.472",
    "mean-versus-median ranks correlated 0.997 and 0.986",
    "median pairwise cosine distances were 0.080, 0.103 and 0.052",
    "median maximum leave-one-slide-out centroid changes were 0.010, 0.012 and 0.006",
    "13/15 continuous and 4/5 binary models",
):
    require(phrase in main_text, f"Pathology, linkage or pooling result is missing: {phrase}")
require("8. Elevate pathology quality control, molecular-slide linkage and pooling limitations"
        in response_text, "The latest pathology-quality reviewer response is missing")

for phrase in (
    "142/95/101 of 243 directly observed genomic-alteration tasks (58.4%/39.1%/41.6%)",
    "95/75/79 of 183 binary composite genomic-context tasks (51.9%/41.0%/43.2%)",
    "30/13/22 of 413 sequencing-derived continuous burdens (7.3%/3.1%/5.3%)",
    "100/47/60 of 279 transcriptomic signatures (35.8%/16.8%/21.5%)",
    "27/18/18 of 638 computationally inferred immune-cell fractions (4.2%/2.8%/2.8%)",
    "40/16/19 of 166 continuous composite scores (24.1%/9.6%/11.4%)",
    "197/94/119 of 1,496 tasks (13.2%/6.3%/8.0%)",
):
    require(phrase in main_text, f"Provenance-stratified main result is missing: {phrase}")
require("9. Separate direct genomic outcomes from derived and same-H&E reference phenotypes"
        in response_text, "The current provenance reviewer response is missing")
for phrase in (
    "We did not perform blinded pathologist review or spatial relevance localisation",
    "they do not contain patch or tile embeddings, attention maps, relevance maps or WSI pixels",
    "nearest-neighbour examples only as qualitative context",
):
    require(phrase in supp_text, f"Spatial-interpretation status is missing: {phrase}")
require("molecular prediction" not in main_text.lower() and
        "molecular-prediction" not in main_text.lower() and
        "immune phenotype prediction" not in main_text.lower(),
        "An unqualified prediction umbrella term remains in the main manuscript")

for phrase in (
    "We obtained immune, inflammatory and genomic-context features from the PanImmune_MS sheet",
    "22 relative immune-cell fractions inferred from bulk RNA sequencing by CIBERSORT with the LM22 reference",
    "We constructed mutation labels from TCGA MC3 calls",
    "We assigned wild type only when MC3 clinical metadata confirmed profiling",
    "We imported pathway status from the Pathway level sheet of Sanchez-Vega et al. Table S4",
    "We obtained the Taylor aneuploidy score",
    "We obtained fusion calls and the assayed denominator",
    "we downloaded MANTIS and MSIsensor fields from the cBioPortal TCGA PanCancer Atlas",
):
    require(phrase in main_text, f"Outcome-acquisition Methods detail is missing: {phrase}")
require("Outcome acquisition, participant linkage and label construction" in supp_text,
        "Supplementary outcome-acquisition Methods are missing")
require("Panel C. Published outcome sources and participant-level construction" in supp_text,
        "Supplementary outcome-source table is missing")
require("Methods section should report how sample-level CIBERSORT, mutation, inflammatory and pathway information was obtained"
        in response_text, "The outcome-acquisition response is missing")
with (ROOT / "data" / "reference" / "outcome_source_acquisition_map.csv").open(
        newline="", encoding="utf-8-sig") as handle:
    outcome_source_rows = list(csv.DictReader(handle))
require(len(outcome_source_rows) == 6 and
        {row["source"] for row in outcome_source_rows} == {
            "Thorsson PanImmune", "MC3 mutations with Bailey driver catalogue",
            "Sanchez-Vega oncogenic pathways", "Taylor aneuploidy",
            "Gao fusions", "cBioPortal TCGA PanCancer MSI",
        }, "Outcome-source acquisition map is incomplete")

with (ROOT / "results" / "tables" /
      "foundation_model_provenance_stratified_summary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    provenance_rows = list(csv.DictReader(handle))
require(len(provenance_rows) == 21,
        "Provenance summary should contain 21 representation-by-class rows")
provenance_by_key = {
    (row["foundation_model"], row["outcome_type"], row["measurement_class"]): row
    for row in provenance_rows
}
expected_provenance = {
    ("TITAN", "binary", "directly observed genomic alteration"): (243, 142),
    ("GigaSSL", "binary", "directly observed genomic alteration"): (243, 95),
    ("ProvGigaPath", "binary", "directly observed genomic alteration"): (243, 101),
    ("TITAN", "continuous", "sequencing-derived continuous burden"): (413, 30),
    ("GigaSSL", "continuous", "sequencing-derived continuous burden"): (413, 13),
    ("ProvGigaPath", "continuous", "sequencing-derived continuous burden"): (413, 22),
}
for key, (eligible, crossing) in expected_provenance.items():
    row = provenance_by_key[key]
    require(int(row["eligible_tasks"]) == eligible and
            int(row["effect_threshold_crossings"]) == crossing,
            f"Provenance summary changed for {key}")

with (ROOT / "results" / "tables" /
      "molecular_slide_linkage_audit.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    linkage_rows = list(csv.DictReader(handle))
require(len(linkage_rows) == 6, "Molecular-slide linkage audit must contain six sources")
linkage_by_source = {row["source"]: row for row in linkage_rows}
require(linkage_by_source["Thorsson2018_PanImmune_MS"]["identifier_resolution"] ==
        "participant identifier only",
        "Thorsson linkage is not labelled participant-only")
for source in (
    "Taylor2018_TableS2", "SanchezVega2018_TableS4", "Gao2018_TableS1",
    "cBioPortal_TCGA_PanCancer", "TCGA_MC3_Bailey2018",
):
    row = linkage_by_source[source]
    require(int(row["exact_slide_sample_patients"]) == int(row["covered_patients"]),
            f"Sample-barcode concordance changed for {source}")

with (ROOT / "results" / "tables" /
      "pathology_qc_no_residual_patient_audit.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    flagged_rows = list(csv.DictReader(handle))
expected_flagged = {
    "TCGA-B6-A0IA", "TCGA-B6-A0WV", "TCGA-55-8203",
    "TCGA-55-8507", "TCGA-22-4596", "TCGA-DX-AB2L",
}
require({row["patient"] for row in flagged_rows} == expected_flagged and
        sum(int(row["narrative_no_residual_tumour_slides"])
            for row in flagged_rows) == 35,
        "Non-adjudicated no-residual patient audit changed")

with (ROOT / "results" / "tables" /
      "nonadjudicated_no_residual_exclusion_all.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    exclusion_rows = list(csv.DictReader(handle))
require(len(exclusion_rows) == 69 and
        sum(row["original_threshold_retained"].upper() == "TRUE"
            for row in exclusion_rows) == 64 and
        sum(row["highlighted_model"].upper() == "TRUE"
            for row in exclusion_rows) == 4 and
        all(row["original_threshold_retained"].upper() == "TRUE" and
            row["material_change_flag"].upper() != "TRUE"
            for row in exclusion_rows
            if row["highlighted_model"].upper() == "TRUE"),
        "Non-adjudicated exclusion sensitivity changed")

with (ROOT / "results" / "tables" /
      "median_pooling_sensitivity_summary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    median_pool = {row["outcome_type"]: row for row in csv.DictReader(handle)}
require(int(median_pool["continuous"]["models"]) == 219 and
        int(median_pool["continuous"]["retained_original_effect_threshold"]) == 216 and
        float(median_pool["continuous"]["correlation_with_mean"]) > 0.99 and
        int(median_pool["binary"]["models"]) == 104 and
        int(median_pool["binary"]["retained_original_effect_threshold"]) == 99 and
        float(median_pool["binary"]["correlation_with_mean"]) > 0.98,
        "Median-pooling sensitivity changed")

with (ROOT / "results" / "tables" /
      "slide_embedding_heterogeneity_summary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    heterogeneity = {row["foundation_model"]: row for row in csv.DictReader(handle)}
require({model: int(heterogeneity[model]["multi_slide_patients"])
         for model in heterogeneity} ==
        {"TITAN": 843, "GigaSSL": 845, "ProvGigaPath": 778},
        "Multi-slide embedding heterogeneity denominators changed")

with (ROOT / "results" / "tables" /
      "sarc_maximum_slide_patient_exclusion_continuous.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    sarc_cont = list(csv.DictReader(handle))
with (ROOT / "results" / "tables" /
      "sarc_maximum_slide_patient_exclusion_binary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    sarc_bin = list(csv.DictReader(handle))
require(len(sarc_cont) == 15 and
        sum(float(row["exclusion_q2"]) >= 0.20 for row in sarc_cont) == 13 and
        len(sarc_bin) == 5 and
        sum(float(row["exclusion_balanced_accuracy"]) >= 0.60
            for row in sarc_bin) == 4,
        "Maximum-slide SARC exclusion sensitivity changed")

with (ROOT / "results" / "tables" /
      "highlighted_model_performance.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    highlighted_rows = list(csv.DictReader(handle))
require(len(highlighted_rows) == 24, "Highlighted table must contain 24 models")
required_repeat_fields = {
    "repeat_metric", "repeat_metric_mean", "repeat_metric_median",
    "repeat_metric_q1", "repeat_metric_q3", "repeat_metric_min",
    "repeat_metric_max", "repeat_metric_sd", "repeat_crossing_threshold",
    "repeat_crossing_count", "repeat_crossing_proportion", "repeat_partitions",
}
require(required_repeat_fields.issubset(highlighted_rows[0]),
        "Highlighted results lack repeat-distribution fields")
require(all(int(row["repeat_partitions"]) == 5 and
            0 <= float(row["repeat_crossing_proportion"]) <= 1 and
            int(row["repeat_crossing_count"]) ==
            round(5 * float(row["repeat_crossing_proportion"]))
            for row in highlighted_rows),
        "Highlighted crossing proportions are inconsistent")

with (ROOT / "results" / "tables" /
      "foundation_model_tss_grouped_sensitivity.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    grouped_rows = list(csv.DictReader(handle))
require(len(grouped_rows) == 1488, "Grouped analysis must contain 1,488 representation-task rows")
infeasible_grouped = [row for row in grouped_rows if row["feasible"].upper() != "TRUE"]
require(len(infeasible_grouped) == 3 and
        {row["foundation_model"] for row in infeasible_grouped} ==
        {"TITAN", "GigaSSL", "ProvGigaPath"} and
        all(row["tumor_type"] == "ACC" and row["endpoint"] == "Genome doubling"
            for row in infeasible_grouped),
        "The expected ACC genome-doubling grouped failures changed")
required_grouped_fields = {
    "grouped_effect", "matched_random_effect", "grouped_minus_matched_effect",
    "n_codes", "realized_outer_folds", "any_single_class_outer_test_fold",
    "any_single_class_inner_validation_fold",
    "any_single_class_inner_training_fold", "grouped_fold_adequacy_flag",
    "grouped_fold_adequacy_reasons", "metric_construction",
}
require(required_grouped_fields.issubset(grouped_rows[0]),
        "Grouped results lack fold-adequacy or matched-random fields")

with (ROOT / "results" / "tables" /
      "foundation_model_tss_grouped_fold_adequacy_summary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    adequacy = {row["outcome_type"]: row for row in csv.DictReader(handle)}
require(int(adequacy["binary"]["tasks"]) == 273 and
        int(adequacy["binary"]["tasks_with_four_outer_folds"]) == 4 and
        int(adequacy["binary"]["tasks_with_single_class_outer_test_fold"]) == 28 and
        int(adequacy["binary"]["tasks_with_single_class_inner_validation_fold"]) == 120 and
        int(adequacy["binary"]["tasks_with_single_class_inner_training_fold"]) == 1 and
        int(adequacy["binary"]["tasks_with_sparse_grouped_folds"]) == 162 and
        int(adequacy["continuous"]["tasks"]) == 223 and
        int(adequacy["continuous"]["tasks_with_four_outer_folds"]) == 6 and
        int(adequacy["continuous"]["tasks_with_sparse_grouped_folds"]) == 51,
        "Grouped-fold adequacy counts changed")

for registry_path in (
    ROOT / "models" / "model_registry.csv",
    ROOT / "models" / "foundation_models" / "model_registry_additions.csv",
):
    with registry_path.open(newline="", encoding="utf-8-sig") as handle:
        registry_rows = list(csv.DictReader(handle))
    for field in (
        "matched_tss_n_codes", "matched_tss_realized_outer_folds",
        "matched_tss_grouped_fold_adequacy_flag",
        "matched_tss_grouped_fold_adequacy_reasons",
        "matched_tss_metric_construction",
    ):
        require(field in registry_rows[0],
                f"Registry {registry_path.name} lacks {field}")

with (ROOT / "results" / "tables" /
      "foundation_model_binary_operating_rule_metrics.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    operating_rows = list(csv.DictReader(handle))
require(len(operating_rows) == 3834, "Operating-rule table must contain 3,834 rows")
for key in ("balanced_accuracy", "sensitivity", "specificity",
            "auroc", "pr_auc", "prevalence", "no_skill_pr_auc"):
    require(all(row.get(key, "") != "" for row in operating_rows),
            f"Operating-rule table contains missing {key}")
# PPV or NPV is mathematically undefined when an operating rule makes no
# positive or no negative calls. Preserve those values as explicit missing
# fields rather than inventing a predictive value.
require(sum(row.get("ppv", "") == "" for row in operating_rows) == 24,
        "Unexpected number of undefined PPV values")
require(sum(row.get("npv", "") == "" for row in operating_rows) == 1,
        "Unexpected number of undefined NPV values")

score_groups: dict[tuple[str, str, str, str], list[dict[str, str]]] = {}
for row in operating_rows:
    key = (row["foundation_model"], row["tumor_type"], row["family"], row["endpoint"])
    score_groups.setdefault(key, []).append(row)
require(len(score_groups) == 1278 and all(len(rows) == 3 for rows in score_groups.values()),
        "Every representation-task row must have exactly three operating rules")
for rows in score_groups.values():
    require(len({row["auroc"] for row in rows}) == 1, "AUROC changed across call rules")
    require(len({row["pr_auc"] for row in rows}) == 1, "PR-AUC changed across call rules")
    require(all(abs(float(row["prevalence"]) - float(row["no_skill_pr_auc"])) < 1e-12
                for row in rows), "No-skill PR-AUC does not equal prevalence")

with (ROOT / "models" / "foundation_models" /
      "model_registry_additions.csv").open(newline="", encoding="utf-8-sig") as handle:
    registry = list(csv.DictReader(handle))
require(len(registry) == 594, "Alternative-representation registry must contain 594 rows")
binary_registry = [row for row in registry if row["outcome_type"] == "binary"]
require(len(binary_registry) == 198, "Alternative binary registry must contain 198 rows")
for row in binary_registry:
    require(row.get("primary_binary_metric") == "outer out-of-fold AUROC",
            "An alternative binary object is not labelled AUROC-centred")
    require(row.get("binary_operating_threshold", "") != "",
            "An alternative binary object lacks its locked training-derived threshold")

with (ROOT / "results" / "tables" /
      "foundation_model_binary_component_range_20_summary.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    component_summary = list(csv.DictReader(handle))
require(len(component_summary) == 3, "Expanded-component summary must contain three rows")
require(all(int(row["tasks"]) == 366 for row in component_summary),
        "Expanded-component summary does not cover 366 tasks per representation")
require(all(int(row["winners_retained"]) == 323 and
            int(row["winner_tasks"]) == 366 for row in component_summary),
        "Expanded-component winner summary is inconsistent")
require(all(int(row["numerical_failures"]) == 0 and
            int(row["fallbacks"]) == 0 for row in component_summary),
        "Expanded-component analysis used a failure or fallback")

with (ROOT / "results" / "tables" /
      "foundation_model_binary_component_range_20_targets.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    component_targets = list(csv.DictReader(handle))
require(len(component_targets) == 366, "Expanded-component target table must have 366 rows")
require(sum(row["winner_retained"].upper() != "TRUE" for row in component_targets) == 43,
        "Expanded-component winner changes must equal 43")

with (ROOT / "results" / "tables" /
      "foundation_model_binary_component_distribution.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    component_distribution = list(csv.DictReader(handle))
distribution_totals: dict[tuple[str, str], int] = {}
for row in component_distribution:
    key = (row["foundation_model"], row["component_grid"])
    distribution_totals[key] = distribution_totals.get(key, 0) + int(row["outer_fits"])
require(len(distribution_totals) == 6 and
        all(value == 1830 for value in distribution_totals.values()),
        "Each representation-grid component distribution must contain 1,830 outer fits")

with (ROOT / "results" / "tables" /
      "foundation_model_binary_component_feature_handling.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    feature_audit = {row["foundation_model"]: row for row in csv.DictReader(handle)}
require({model: int(feature_audit[model]["constant_features"])
         for model in feature_audit} ==
        {"TITAN": 0, "GigaSSL": 5, "ProvGigaPath": 0},
        "Constant-feature counts are inconsistent")
require(all(int(row["near_constant_features"]) == 0 for row in feature_audit.values()),
        "Unexpected near-constant feature was reported")

with (ROOT / "results" / "tables" /
      "foundation_model_TITAN_ProvGigaPath_binary_component_range.csv").open(
          newline="", encoding="utf-8-sig") as handle:
    titan_prov_rows = list(csv.DictReader(handle))
require(len(titan_prov_rows) == 1 and
        int(titan_prov_rows[0]["pairwise_leader_changed"]) == 30,
        "TITAN versus Prov-GigaPath paired summary is inconsistent")

check_log = ROOT.parent / "PathoFMPred.Rcheck" / "00check.log"
require(check_log.is_file(), "R CMD check log is missing")
check_text = check_log.read_text(encoding="utf-8", errors="replace")
require("Status: 1 NOTE" in check_text and "ERROR" not in check_text and "WARNING" not in check_text,
        "R CMD check did not finish with only the expected new-submission NOTE")

for phrase in (
    "cross-validated coefficient of determination (Q²)",
    "out-of-fold (OOF)",
    "the q-value is the false-discovery-rate-adjusted empirical p-value",
    "selection-conditioned (SC) patient-resampling interval for repeated OOF predictions",
    "TCGA tissue-source-site codes, the barcode-derived submitting-centre fields",
):
    require(phrase in main_text, f"First-use definition is missing: {phrase}")
for phrase in (
    "OOF, out-of-fold",
    "q-value, false-discovery-rate-adjusted p-value",
    "SC interval, selection-conditioned patient-resampling interval",
    "Q², cross-validated coefficient of determination",
    "TCGA tissue-source-site code, barcode-derived submitting-centre field",
):
    require(phrase in main_text, f"Abbreviation-list definition is missing: {phrase}")
require("Initial primary screen; five-repeat mean, SC interval and partition distribution"
        in supp_text, "Initial-screen and five-repeat estimates are not co-labelled")
require("sample-size maturity stratum" in combined.lower(),
        "Sample-size maturity stratum terminology is missing")
require("standard evidence" not in combined.lower() and
        "standard-evidence" not in combined.lower(),
        "Obsolete standard-evidence terminology remains")
require("Screening tier A and tier B are retained only as database-navigation tags"
        in supp_text, "Threshold categories are not limited to navigation")
require("targeted narrative literature audit" in combined.lower(),
        "The literature crosswalk is not labelled as a targeted narrative audit")

main_results = main_text.split("\nResults\n", 1)[1].split("\nDiscussion\n", 1)[0]
for operational_term in ("SHA-256", "checksum", "score-rank construction", "model hash"):
    require(operational_term.lower() not in main_results.lower(),
            f"Implementation detail remains in the biological Results: {operational_term}")
require("standardized LDA score" not in combined and "standardised LDA score" not in combined,
        "A standardized LDA-score plot remains without a verified threshold")

for html_name in ("COAD_example_A.html", "COAD_example_B.html"):
    report_text = (ROOT / "results" / "reports" / html_name).read_text(
        encoding="utf-8", errors="replace"
    )
    require("Reference rank, not probability" in report_text,
            f"{html_name} lacks the immediate rank warning")
    require("Treatment and response narratives" not in report_text,
            f"{html_name} contains treatment-response narration")

for phrase in (
    "reproducible translational research prioritisation resource",
    "does not claim clinical validation",
    "Medical Bioinformatics, Molecular Pathology, Disease Biomarkers and Translational Imaging",
):
    require(phrase in cover_text, f"Cover-letter positioning is missing: {phrase}")
require("Minor and editorial comments" in response_text,
        "The minor/editorial point-by-point response is missing")

print("AUROC-centred release audit passed")
print("Main figures: 3; main tables: 3; references: 42; supplementary figures: 7")
print("Operating-rule rows: 3,834; representation-task score groups: 1,278")
print("PathoFMPred registry rows: 917; R CMD check status: 1 NOTE")
