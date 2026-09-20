#!/usr/bin/env python3

"""Compact Figure 1 so the patient-overlap revision does not add a page."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


DOCX = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/manuscript/"
    "manuscript_JTM_multifoundation_atlas.docx"
)
FIGURE = Path(
    "/Users/stefano/Documents/Titan/titan-prediction/results/figures/"
    "Figure1_patient_overlap_venn.png"
)


def starting(document, prefix):
    return next(p for p in document.paragraphs if p.text.startswith(prefix))


def replace_text(paragraph, text):
    for child in list(paragraph._p):
        if not child.tag.endswith("}pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


doc = Document(DOCX)

cohort = starting(doc, "The released artifacts comprised 11,449 TITAN slides")
replace_text(
    cohort,
    "The released artifacts comprised 11,449 TITAN slides from 9,404 patients, "
    "11,427 Giga-SSL slides from 9,378 patients and 10,328 Prov-GigaPath slides "
    "from 8,393 patients. Their union contained 9,471 unique patients. Of these, "
    "8,241 (87.0%) occurred in all three datasets and formed the matched benchmark; "
    "Figure 1 shows the pairwise and dataset-specific counts.",
)

caption = starting(doc, "Figure 1. Patient overlap among ")
replace_text(
    caption,
    "Figure 1. Patient overlap among TITAN, Giga-SSL and Prov-GigaPath. Values "
    "are mutually exclusive patient counts; the 8,241-patient three-way "
    "intersection formed the matched benchmark. Circle areas are schematic.",
)

image_paragraph = caption._p.getprevious()
image_wrapper = next(p for p in doc.paragraphs if p._p is image_paragraph)
for child in list(image_wrapper._p):
    if not child.tag.endswith("}pPr"):
        image_wrapper._p.remove(child)
image_wrapper.alignment = WD_ALIGN_PARAGRAPH.CENTER
shape = image_wrapper.add_run().add_picture(str(FIGURE), width=Inches(6.0))
shape._inline.docPr.set(
    "descr",
    "Non-area-proportional Venn diagram of patient overlap among TITAN, "
    "Giga-SSL and Prov-GigaPath TCGA embedding datasets.",
)
shape._inline.docPr.set("title", "Patient overlap across embedding datasets")

pooling = starting(
    doc,
    "Within each representation, eligible primary-tumour diagnostic slides were pooled",
)
cohort._p.addnext(pooling._p)

doc.save(DOCX)
print(DOCX)
