#!/usr/bin/env python3
"""Remove the dedicated TITAN-only Results subsection from the main manuscript."""

from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
BACKUP = DOCX.with_name(
    "manuscript_JTM_multifoundation_atlas.before_titan_results_subsection_removal.docx"
)
START = "Supporting TITAN screen and translational priorities"
END = "Discussion"
PARTITION_PREFIX = "Representation ordering depended on the downstream analysis."
NOTE = (
    " The separate TITAN-only permutation/FDR screen is reported in the Supplement and "
    "is not used to rank the three representations."
)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


copy2(DOCX, BACKUP)
document = Document(DOCX)
paragraphs = list(document.paragraphs)

starts = [i for i, paragraph in enumerate(paragraphs) if paragraph.text.strip() == START]
if len(starts) != 1:
    raise RuntimeError(f"Expected one TITAN subsection heading, found {len(starts)}")
start = starts[0]

ends = [
    i
    for i, paragraph in enumerate(paragraphs[start + 1 :], start + 1)
    if paragraph.text.strip() == END
]
if len(ends) != 1:
    raise RuntimeError(f"Expected one following Discussion heading, found {len(ends)}")
end = ends[0]

partition = [
    paragraph for paragraph in paragraphs if paragraph.text.startswith(PARTITION_PREFIX)
]
if len(partition) != 1:
    raise RuntimeError(f"Expected one partition paragraph, found {len(partition)}")
if NOTE.strip() not in partition[0].text:
    partition[0].add_run(NOTE)

for paragraph in paragraphs[start:end]:
    delete_paragraph(paragraph)

document.save(DOCX)
print(f"Removed {end - start} paragraphs from {DOCX}")
