from pathlib import Path
import re
import shutil

from docx import Document


DOCX = Path('/Users/stefano/Documents/Titan/titan-prediction/manuscript/manuscript_JTM_multifoundation_atlas.docx')
BACKUP = DOCX.with_name('manuscript_JTM_multifoundation_atlas.before_reference_reorder.docx')


def find_paragraph(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f'Expected one paragraph beginning {prefix!r}; found {len(matches)}')
    return matches[0]


def replace_once(paragraph, old, new):
    if old not in paragraph.text:
        raise RuntimeError(f'Text not found in paragraph: {old!r}')
    if len(paragraph.runs) != 1:
        raise RuntimeError(f'Expected one run in target paragraph, found {len(paragraph.runs)}')
    paragraph.runs[0].text = paragraph.text.replace(old, new, 1)


def append_text(paragraph, text):
    if len(paragraph.runs) != 1:
        raise RuntimeError(f'Expected one run in target paragraph, found {len(paragraph.runs)}')
    paragraph.runs[0].text = paragraph.text.rstrip() + text


def parse_group(raw):
    numbers = []
    for token in raw.split(','):
        token = token.strip()
        if '-' in token:
            start, end = (int(x.strip()) for x in token.split('-', 1))
            numbers.extend(range(start, end + 1))
        else:
            numbers.append(int(token))
    return numbers


def citation_ids(text):
    ids = []
    for match in re.finditer(r'\[([0-9][0-9,\- ]*)\]', text):
        ids.extend(parse_group(match.group(1)))
    return ids


def remap_citations(text, mapping):
    def repl(match):
        old_ids = parse_group(match.group(1))
        new_ids = []
        for old_id in old_ids:
            if old_id not in mapping:
                raise RuntimeError(f'Citation {old_id} has no remapping')
            mapped = mapping[old_id]
            if mapped not in new_ids:
                new_ids.append(mapped)
        new_ids.sort()
        return '[' + ','.join(str(x) for x in new_ids) + ']'
    return re.sub(r'\[([0-9][0-9,\- ]*)\]', repl, text)


if not BACKUP.exists():
    shutil.copy2(DOCX, BACKUP)
doc = Document(BACKUP)

# Add citations at their first relevant scientific or resource mention.
p = find_paragraph(doc, 'The narrower unresolved gap is a reproducible multi-representation atlas')
replace_once(p, 'PathoFMPred is a reproducible analysis interface', 'PathoFMPred [32] is a reproducible analysis interface')

p = find_paragraph(doc, 'We evaluated three released embedding pipelines')
replace_once(p, 'TCGA embeddings [39]', 'TCGA embeddings [39,41]')
replace_once(p, 'public TCGA embedding dataset [40]', 'public TCGA embedding dataset [40,42]')

p = find_paragraph(doc, 'Outcomes were analysed within cancer and classified as directly observed genomic alterations')
append_text(
    p,
    ' Source resources included the TCGA immune landscape [19], aneuploidy scores [20], driver fusions [21], microsatellite-instability calls [22], cBioPortal [23], MC3 mutation calls [24], driver-gene annotation [25], oncogenic-pathway scores [26] and the TCGA clinical resource [29]. CIBERSORT fractions [37] and the H&E-derived TIL fraction [38] were retained as computational reference phenotypes rather than direct immune-cell measurements.'
)

p = find_paragraph(doc, 'Within each cancer–endpoint pair, nested patient-level five-fold cross-validation')
append_text(p, ' The PLS and double-cross-validation framework followed established methodological references [17,18].')

p = find_paragraph(doc, 'The supporting TITAN-only screen evaluated 2,073 eligible pairs')
replace_once(p, 'Benjamini–Hochberg correction was applied', 'Benjamini–Hochberg correction [27] was applied')

p = find_paragraph(doc, 'This retrospective computational benchmark was not prospectively registered.')
replace_once(
    p,
    'analysis_chronology.csv.',
    'analysis_chronology.csv in the companion repository [30]. Reporting was mapped to TRIPOD+AI [28].'
)

p = find_paragraph(doc, 'Analyses used R 4.6.0 and fastPLS 0.3.')
replace_once(p, 'The companion repository contains', 'The companion repository [30] contains')

p = find_paragraph(doc, 'The endpoint novelty is moderate.')
append_text(
    p,
    ' The established association between GTF2I mutation and spindle-cell thymoma morphology was considered separately from predictive-model novelty [35].'
)

p = find_paragraph(doc, 'The analysis used the official TITAN TCGA feature artifact')
replace_once(
    p,
    'The analysis used the official TITAN TCGA feature artifact, official Giga-SSL TCGA embeddings and the seandavis/tcga_provgigapath_embeddings Hugging Face dataset.',
    'The analysis used the official TITAN TCGA feature artifact [16], official Giga-SSL TCGA embeddings [41] and the seandavis/tcga_provgigapath_embeddings Hugging Face dataset [42]. Participant and molecular metadata were obtained from the Genomic Data Commons and the cited TCGA source studies [33].'
)
replace_once(p, 'companion repository [30]', 'companion repository [30]')
replace_once(p, 'PathoFMPred source remains private', 'PathoFMPred source [32] remains private')

# Identify reference paragraphs and parse the original bibliography.
ref_heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == 'References')
ref_paragraphs = []
old_references = {}
for p in doc.paragraphs[ref_heading_index + 1:]:
    match = re.match(r'^(\d+)\.\s+(.*)$', p.text.strip(), flags=re.S)
    if not match:
        continue
    old_number = int(match.group(1))
    old_references[old_number] = match.group(2)
    ref_paragraphs.append(p)

if sorted(old_references) != list(range(1, 43)):
    raise RuntimeError(f'Expected original references 1 to 42; found {sorted(old_references)}')
if len(ref_paragraphs) != 42:
    raise RuntimeError(f'Expected 42 reference paragraphs; found {len(ref_paragraphs)}')

# Determine order of first appearance in the main text.
first_appearance = []
for p in doc.paragraphs[:ref_heading_index]:
    for old_id in citation_ids(p.text):
        if old_id not in first_appearance:
            first_appearance.append(old_id)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for old_id in citation_ids(p.text):
                    if old_id not in first_appearance:
                        first_appearance.append(old_id)

missing = sorted(set(old_references) - set(first_appearance))
unknown = sorted(set(first_appearance) - set(old_references))
if missing != [34, 36] or unknown:
    raise RuntimeError(f'Reference audit failed. Uncited={missing}; unknown citations={unknown}')

mapping = {old_id: new_id for new_id, old_id in enumerate(first_appearance, start=1)}

# Renumber every numeric citation in body paragraphs and tables.
for p in doc.paragraphs[:ref_heading_index]:
    updated = remap_citations(p.text, mapping)
    if updated != p.text:
        if len(p.runs) != 1:
            raise RuntimeError(f'Cannot safely renumber multi-run paragraph: {p.text[:100]}')
        p.runs[0].text = updated
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                updated = remap_citations(p.text, mapping)
                if updated != p.text:
                    if len(p.runs) != 1:
                        raise RuntimeError(f'Cannot safely renumber multi-run table paragraph: {p.text[:100]}')
                    p.runs[0].text = updated

# Rebuild the numbered bibliography in first-appearance order and remove the
# two entries that do not support a claim retained in the main manuscript.
for new_number, (p, old_number) in enumerate(zip(ref_paragraphs, first_appearance), start=1):
    new_text = f'{new_number}. {old_references[old_number]}'
    if len(p.runs) == 1:
        p.runs[0].text = new_text
    else:
        p.text = new_text
for p in ref_paragraphs[len(first_appearance):]:
    p._element.getparent().remove(p._element)

# Final structural checks.
all_final_citations = []
for p in doc.paragraphs[:ref_heading_index]:
    all_final_citations.extend(citation_ids(p.text))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                all_final_citations.extend(citation_ids(p.text))

first_final = []
for final_id in all_final_citations:
    if final_id not in first_final:
        first_final.append(final_id)
expected_final = list(range(1, len(first_appearance) + 1))
if first_final != expected_final:
    raise RuntimeError(f'First-appearance order is not sequential: {first_final}')
if set(all_final_citations) != set(expected_final):
    raise RuntimeError('Not every bibliography entry is cited after renumbering')

doc.save(DOCX)

print(f'Updated: {DOCX}')
print(f'Backup:  {BACKUP}')
print('Old to new reference mapping:')
print(' '.join(f'{old}->{new}' for old, new in sorted(mapping.items(), key=lambda x: x[1])))
