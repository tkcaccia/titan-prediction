#!/usr/bin/env python3
"""Define cancer-endpoint pairs and state how they are analysed."""

from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript" / "manuscript_JTM_multifoundation_atlas.docx"
BACKUP = DOCX.with_name(
    "manuscript_JTM_multifoundation_atlas.before_cancer_endpoint_clarification.docx"
)

ABSTRACT_OLD = (
    "Methods. The retrospective matched benchmark compared TITAN, Giga-SSL and Prov-GigaPath "
    "on the same 8,241 patients and 1,933 cancer-endpoint pairs using identical patient-level "
    "folds and one 1-10-component PLS regression/PLS-LDA probe. Slides were pooled before "
    "outcome matching."
)
ABSTRACT_NEW = (
    "Methods. The retrospective matched benchmark compared TITAN, Giga-SSL and Prov-GigaPath "
    "on the same 8,241 patients and 1,933 cancer-specific prediction tasks, each defined by one "
    "TCGA cancer type and one endpoint. Slides were pooled before outcome matching, and every "
    "task was analysed only within its cancer using identical patient-level folds and one "
    "1-10-component PLS regression/PLS-LDA probe across the three representations."
)

BACKGROUND_ANCHOR = (
    "TITAN is a multimodal whole-slide model trained by visual self-supervision and "
    "vision-language alignment"
)
BACKGROUND_DEFINITION = (
    "In this study, a cancer-endpoint pair denotes one TCGA cancer type combined with one "
    "outcome, for example COAD with APC mutation or LIHC with the wound-healing score. It "
    "defines a cancer-specific prediction task, not a pair of patients. For every eligible "
    "pair, only patients from that cancer with the required outcome label were analysed. The "
    "pair was fitted separately to each available representation; in the matched benchmark, "
    "TITAN, Giga-SSL and Prov-GigaPath used identical patients and validation folds. No "
    "prediction model pooled different cancer types."
)

SUMMARY_OLD = (
    "Using identical patients, folds, seeds and tuning rules, we evaluated 1,933 shared "
    "cancer-endpoint pairs with nested patient-level PLS regression or PLS-LDA."
)
SUMMARY_NEW = (
    "Using identical patients, folds, seeds and tuning rules, we evaluated 1,933 such tasks "
    "with nested patient-level PLS regression or PLS-LDA."
)

MODELS_OLD = (
    "Within each cancer-endpoint pair, nested patient-level five-fold cross-validation "
    "centred features within training folds"
)
MODELS_NEW = (
    "Each eligible cancer-endpoint pair was analysed as a separate cancer-specific task. In "
    "the matched benchmark, the TITAN, Giga-SSL and Prov-GigaPath models were fitted "
    "independently but used the same patients and validation folds. Nested patient-level "
    "five-fold cross-validation centred features within training folds"
)


def replace_once(document, old, new):
    matches = [paragraph for paragraph in document.paragraphs if old in paragraph.text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph containing {old!r}, found {len(matches)}")
    paragraph = matches[0]
    revised = paragraph.text.replace(old, new)
    first = paragraph.runs[0]
    first.text = revised
    for run in paragraph.runs[1:]:
        run.text = ""
    return paragraph


def insert_after(paragraph, text):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.style = paragraph.style
    inserted.add_run(text)
    return inserted


def add_terminology_row(document):
    tables = [
        table
        for table in document.tables
        if len(table.rows) > 1
        and table.cell(0, 0).text.strip() == "Term"
        and table.cell(0, 1).text.strip() == "Meaning in this study"
    ]
    if len(tables) != 1:
        raise RuntimeError(f"Expected one terminology table, found {len(tables)}")
    table = tables[0]
    existing = [row.cells[0].text.strip() for row in table.rows]
    if "Cancer-endpoint pair" in existing:
        raise RuntimeError("Cancer-endpoint pair row already exists")

    row = table.add_row()
    row.cells[0].text = "Cancer-endpoint pair"
    row.cells[1].text = (
        "One TCGA cancer type combined with one outcome. It defines one cancer-specific "
        "prediction task; each representation is fitted separately, and no model pools "
        "cancer types."
    )
    table.rows[0]._tr.addnext(row._tr)

    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(7.4)


copy2(DOCX, BACKUP)
document = Document(DOCX)

replace_once(document, ABSTRACT_OLD, ABSTRACT_NEW)

anchors = [p for p in document.paragraphs if BACKGROUND_ANCHOR in p.text]
if len(anchors) != 1:
    raise RuntimeError(f"Expected one background anchor, found {len(anchors)}")
insert_after(anchors[0], BACKGROUND_DEFINITION)

replace_once(document, SUMMARY_OLD, SUMMARY_NEW)
replace_once(document, MODELS_OLD, MODELS_NEW)
add_terminology_row(document)

document.save(DOCX)
print(f"Clarified cancer-endpoint pairs in {DOCX}")
